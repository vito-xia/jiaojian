import io
import json
import os
import shutil
import subprocess
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from process_data import read_controls, read_delivery_monitor, read_scores
from 数据源.脚本.平台管控积分 import update_platform_sources as sync


CONTROL_HEADERS = (
    "管控单id", "网点编码", "网点名称", "快递公司", "所在地区", "商家id", "商家名称",
    "违规场景", "管控状态", "管控动作", "开始日期", "预计完结日期", "管控机制", "管控类型",
)
SCORE_HEADERS = [
    "网点名称", "网点编码", "网点所在省份", "网点所在城市", "网点所在区县", "网点所在街道",
    "快递商名称", "涉及商家店铺ID", "涉及商家名称", "违规场景", "异常等级", "协同状态",
    "物流商反馈一级异常原因", "物流商反馈二级异常原因", "当前违规积分", "累计违规积分", "违规日期",
] + [f"指标{i}" for i in range(17, 44)]
for index, name in {
    20: "发运超时", 21: "发运超时异常单量", 22: "发运超时操作单量",
    38: "派签超时", 39: "派签超时异常单量", 40: "派签超时操作单量",
    41: "超时未派", 42: "超时未派异常单量", 43: "超时未派操作单量",
}.items():
    SCORE_HEADERS[index] = name


def control_row(identifier, day, delivery=False, status="执行中", scene=None):
    day = datetime.fromisoformat(day) if len(day) == 10 else day
    return [identifier, "000001", "网点甲", "韵达", "省区", None, None,
            scene or (sync.DELIVERY_SCENE if delivery else "物流停滞-揽收端"), status,
            "限制面单到达" if delivery else "限制面单取号", day, datetime(2026, 10, 31), "积分制", "网点级"]


def score_row(day, points, delivery=False, merchant="全部商家", branch="网点甲"):
    row = [None] * len(SCORE_HEADERS)
    for index, value in {
        0: branch, 1: "000001", 6: "韵达", 8: merchant,
        9: sync.DELIVERY_SCENE if delivery else "物流停滞-揽收端",
        10: "普通异常", 11: "待反馈", 14: points, 16: day,
        20: "0%", 21: 0, 22: None, 38: "0%", 39: 0, 40: None, 41: None, 42: None, 43: 0,
    }.items():
        row[index] = value
    return row


def save_book(path, tables):
    book = Workbook()
    book.remove(book.active)
    for name, headers, rows in tables:
        sheet = book.create_sheet(name)
        sheet.append(list(headers))
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "B2"
        sheet.column_dimensions["A"].width = 29
        sheet.row_dimensions[1].height = 35
        sheet.auto_filter.ref = sheet.dimensions
        if rows:
            sheet.row_dimensions[2].height = 27
            sheet["A2"].fill = PatternFill("solid", fgColor="FFFF0000")
            sheet["A2"].comment = Comment("历史备注", "User")
        if "开始日期" in headers:
            for row in sheet.iter_rows(min_row=2):
                row[10].number_format = "mm-dd-yy"
                row[11].number_format = "mm-dd-yy"
    book.save(path)
    book.close()


def rows_of(path, sheet_name):
    book = load_workbook(path, data_only=False)
    try:
        return [list(row) for row in book[sheet_name].iter_rows(min_row=2, values_only=True)]
    finally:
        book.close()


def fixture(root):
    data_dir = root / "数据源"
    manual = data_dir / "数据源-手动更新"
    manual.mkdir(parents=True)
    control_dir = manual / "系统导出-管控"
    score_dir = manual / "系统管控-积分"
    control_dir.mkdir()
    score_dir.mkdir()
    paths = {key: manual / name for key, (name, _) in sync.TARGETS.items()}
    save_book(paths["④"], [("异常信息", SCORE_HEADERS, [score_row("2026-10-05", 8), score_row("2026-10-06", 9)])])
    save_book(paths["⑤"], [("Sheet1", CONTROL_HEADERS,
                            [control_row("P-old", "2026-10-01"), control_row("P6", "2026-10-06")])])
    save_book(paths["⑥"], [
        ("每日派送管控", CONTROL_HEADERS,
         [control_row("D-old", "2026-10-02", True), control_row("D6", "2026-10-06", True)]),
        ("每日派送积分", SCORE_HEADERS, [score_row("2026-10-05", 8, True), score_row("2026-10-06", 9, True)]),
    ])
    control_export = control_dir / "管控.xlsx"
    score_export = score_dir / "积分.xlsx"
    save_book(control_export, [("Sheet1", CONTROL_HEADERS, [
        control_row("D7", "2026-10-07 03:15:06", scene="物流停滞-揽收端,物流停滞-派送端"),
        control_row("P7", "2026-10-07 02:00:01", status="执行中-预警期"),
        control_row("P6", "2026-10-06 09:30:00", status="已完结"),
    ])])
    save_book(score_export, [("异常信息", SCORE_HEADERS, [
        score_row("2026-10-07", 2), score_row("2026-10-06", 0),
        score_row("2026-10-07", 3, True), score_row("2026-10-06", 1, True),
        score_row("无效日期也不参与筛选", 99, merchant="单个商家"),
    ])])
    return data_dir, paths, control_export, score_export


def update(data_dir):
    plan = sync.plan_sources(data_dir)
    try:
        return sync.commit_plan(plan, data_dir / "备份")
    finally:
        plan.close()


class PlatformSourceTests(unittest.TestCase):
    def test_overlap_new_day_split_dates_status_styles_and_native_readers(self):
        with tempfile.TemporaryDirectory() as temp:
            data, paths, control_source, score_source = fixture(Path(temp))
            source_bytes = (control_source.read_bytes(), score_source.read_bytes())
            backup = update(data)
            platform = rows_of(paths["⑤"], "Sheet1")
            delivery = rows_of(paths["⑥"], "每日派送管控")
            self.assertEqual([r[0] for r in platform], ["P-old", "P6", "P7"])
            self.assertEqual([r[8] for r in platform], ["已完结", "已完结", "执行中-预警期"])
            self.assertEqual([r[0] for r in delivery], ["D-old", "D6", "D7"])
            self.assertEqual([r[8] for r in delivery], ["已完结", "已完结", "执行中"])
            self.assertTrue(all(isinstance(r[10], datetime) and r[10].hour == 0 for r in platform + delivery))
            for path, name in [(paths["④"], "异常信息"), (paths["⑤"], "Sheet1"), (paths["⑥"], "每日派送积分")]:
                book = load_workbook(path)
                sheet = book[name]
                self.assertEqual(sheet.freeze_panes, "B2")
                self.assertEqual(sheet.column_dimensions["A"].width, 29)
                self.assertEqual(sheet["A2"].fill.fgColor.rgb, "FFFF0000")
                self.assertEqual(sheet["A2"].comment.text, "历史备注")
                self.assertEqual(sheet.row_dimensions[2].height, 27)
                self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
                book.close()
            platform_scores, daily, _ = read_scores(paths["④"].parent)
            self.assertEqual([r["current_score"] for r in platform_scores], [8, 0, 2])
            self.assertEqual(daily["网点甲"]["2026-10-06"], 0)
            self.assertEqual(platform_scores[-1]["shipment_timeout_abnormal_count"], 0)
            self.assertIsNone(platform_scores[-1]["shipment_timeout_operation_count"])
            controls = read_controls(paths["⑤"].parent)
            self.assertEqual(controls[-1]["start_date"], "2026-10-07")
            d_controls, d_scores, d_daily, _, _ = read_delivery_monitor(paths["⑥"].parent)
            self.assertEqual(len(d_controls), 3)
            self.assertEqual([r["current_score"] for r in d_scores], [8, 1, 3])
            self.assertEqual(d_daily["网点甲"]["2026-10-07"], 3)
            self.assertEqual(rows_of(paths["④"], "异常信息")[-1][1], "000001")
            self.assertEqual(source_bytes, (control_source.read_bytes(), score_source.read_bytes()))
            report = json.loads((backup / "更新报告.json").read_text(encoding="utf-8"))
            self.assertEqual(report["status"], "completed")
            self.assertEqual(len(list(backup.glob("*.xlsx"))), 3)

    def test_repeat_run_has_same_rows_and_no_duplicate_append(self):
        with tempfile.TemporaryDirectory() as temp:
            data, paths, _, _ = fixture(Path(temp))
            update(data)
            first = {(p.name, s): rows_of(p, s) for p in paths.values() for s in sync.TARGETS[p.name[0]][1]}
            update(data)
            second = {(p.name, s): rows_of(p, s) for p in paths.values() for s in sync.TARGETS[p.name[0]][1]}
            self.assertEqual(first, second)

    def test_missing_score_scene_clears_covered_days_but_keeps_history(self):
        with tempfile.TemporaryDirectory() as temp:
            data, paths, _, source = fixture(Path(temp))
            save_book(source, [("异常信息", SCORE_HEADERS, [score_row("2026-10-06", 0), score_row("2026-10-07", 2)])])
            update(data)
            self.assertEqual([r[16] for r in rows_of(paths["⑥"], "每日派送积分")], ["2026-10-05"])

    def test_control_scene_without_new_rows_preserves_history_and_completes_it(self):
        with tempfile.TemporaryDirectory() as temp:
            data, paths, source, _ = fixture(Path(temp))
            save_book(source, [("Sheet1", CONTROL_HEADERS, [control_row("P6", "2026-10-06")])])
            update(data)
            rows = rows_of(paths["⑥"], "每日派送管控")
            self.assertEqual([r[0] for r in rows], ["D-old", "D6"])
            self.assertEqual([r[8] for r in rows], ["已完结", "已完结"])

    def test_latest_file_and_column_mapping_ignore_lock_and_wrong_dimensions(self):
        with tempfile.TemporaryDirectory() as temp:
            data, paths, source, _ = fixture(Path(temp))
            old = source.parent / "旧导出.xlsx"
            old.write_bytes(source.read_bytes())
            os.utime(old, ns=(1, 1))
            (source.parent / "~$最新临时.xlsx").write_bytes(b"not an xlsx")
            rows = rows_of(source, "Sheet1")
            order = list(reversed(range(len(CONTROL_HEADERS))))
            save_book(source, [("Sheet1", [CONTROL_HEADERS[i] for i in order], [[r[i] for i in order] for r in rows])])
            output = io.BytesIO()
            with ZipFile(source) as original, ZipFile(output, "w") as rewritten:
                for info in original.infolist():
                    content = original.read(info.filename)
                    if info.filename == "xl/worksheets/sheet1.xml":
                        content = content.replace(b'<dimension ref="A1:N4"/>', b'<dimension ref="A1:A1"/>')
                    rewritten.writestr(info, content)
            source.write_bytes(output.getvalue())
            plan = sync.plan_sources(data)
            try:
                self.assertEqual(plan.exports["管控"], source)
                self.assertEqual(sum(u.summary["incoming_rows"] for u in plan.updates if u.summary["kind"] == "管控"), 3)
                sync.commit_plan(plan, data / "备份")
            finally:
                plan.close()
            self.assertEqual(rows_of(paths["⑤"], "Sheet1")[-1][0], "P7")

    def test_check_is_read_only(self):
        with tempfile.TemporaryDirectory() as temp:
            data, _, _, _ = fixture(Path(temp))
            before = {str(p.relative_to(data)): p.read_bytes() for p in data.rglob("*") if p.is_file()}
            with patch.object(sync.sys, "argv", ["update", "--data-dir", str(data), "--check"]), patch("builtins.print"):
                self.assertEqual(sync.main(), 0)
            after = {str(p.relative_to(data)): p.read_bytes() for p in data.rglob("*") if p.is_file()}
            self.assertEqual(before, after)

    def test_missing_or_invalid_exports_fail_before_writing(self):
        for case in ("missing", "missing_header", "invalid_date", "missing_last_day", "empty", "duplicate_id"):
            with self.subTest(case=case), tempfile.TemporaryDirectory() as temp:
                data, paths, control, score = fixture(Path(temp))
                before = {p: p.read_bytes() for p in paths.values()}
                if case == "missing":
                    score.unlink()
                elif case == "missing_header":
                    latest = score.parent / "损坏的新导出.xlsx"
                    save_book(latest, [("异常信息", ["错误列"], [[1]])])
                    os.utime(latest, ns=(score.stat().st_mtime_ns + 10**9,) * 2)
                elif case == "invalid_date":
                    save_book(control, [("Sheet1", CONTROL_HEADERS, [control_row("bad", "不是日期")])])
                elif case == "missing_last_day":
                    save_book(score, [("异常信息", SCORE_HEADERS, [score_row("2026-10-07", 2)])])
                elif case == "empty":
                    save_book(score, [("异常信息", SCORE_HEADERS, [])])
                else:
                    save_book(control, [("Sheet1", CONTROL_HEADERS, [control_row("same", "2026-10-06")] * 2)])
                with self.assertRaises(sync.SourceUpdateError):
                    sync.plan_sources(data)
                self.assertEqual(before, {p: p.read_bytes() for p in paths.values()})
                self.assertFalse((data / "备份").exists())

    def test_second_replacement_failure_restores_original_bytes(self):
        with tempfile.TemporaryDirectory() as temp:
            data, paths, _, _ = fixture(Path(temp))
            before = {p: p.read_bytes() for p in paths.values()}
            real_replace = sync.os.replace
            def fail_second(src, dst):
                if Path(dst) == paths["⑤"]:
                    raise PermissionError("模拟 Excel 占用")
                return real_replace(src, dst)
            with patch.object(sync.os, "replace", side_effect=fail_second):
                with self.assertRaisesRegex(sync.SourceUpdateError, "已恢复"):
                    update(data)
            self.assertEqual(before, {p: p.read_bytes() for p in paths.values()})
            report = next((data / "备份").glob("*/更新报告.json"))
            self.assertEqual(json.loads(report.read_text(encoding="utf-8"))["status"], "rolled_back")
            self.assertFalse(list(paths["④"].parent.glob(".platform-*")))

    def test_staged_validation_failure_leaves_targets_untouched(self):
        with tempfile.TemporaryDirectory() as temp:
            data, paths, _, _ = fixture(Path(temp))
            before = {p: p.read_bytes() for p in paths.values()}
            with patch.object(sync, "validate_staged", side_effect=sync.SourceUpdateError("模拟保存校验失败")):
                with self.assertRaisesRegex(sync.SourceUpdateError, "保存校验失败"):
                    update(data)
            self.assertEqual(before, {p: p.read_bytes() for p in paths.values()})
            self.assertFalse((data / "备份").exists())

    def test_file_changed_after_planning_stops_replacement(self):
        with tempfile.TemporaryDirectory() as temp:
            data, paths, source, _ = fixture(Path(temp))
            before = {p: p.read_bytes() for p in paths.values()}
            plan = sync.plan_sources(data)
            try:
                os.utime(source, ns=(source.stat().st_atime_ns, source.stat().st_mtime_ns + 10**9))
                with self.assertRaisesRegex(sync.SourceUpdateError, "发生变化"):
                    sync.commit_plan(plan, data / "备份")
            finally:
                plan.close()
            self.assertEqual(before, {p: p.read_bytes() for p in paths.values()})

    @unittest.skipUnless(os.name == "nt", "Windows BAT integration")
    def test_root_entry_stops_before_dashboard_on_source_failure(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            entry = root / "更新看板数据.bat"
            shutil.copy2(sync.PROJECT_ROOT / entry.name, entry)
            script = root / "数据源" / "脚本" / "平台管控积分" / "update_platform_sources.py"
            script.parent.mkdir(parents=True)
            shutil.copy2(Path(sync.__file__), script)
            marker = root / "should_not_run.txt"
            (root / "process_data.py").write_text("from pathlib import Path\nPath('should_not_run.txt').write_text('ran')\n", encoding="utf-8")
            result = subprocess.run(["cmd.exe", "/d", "/c", str(entry), "--no-build", "--no-pause"],
                                    cwd=root, capture_output=True, timeout=30)
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse(marker.exists())
            self.assertEqual(result.stderr.decode("utf-8", errors="replace").count("积分管控更新失败"), 1)


if __name__ == "__main__":
    unittest.main()
