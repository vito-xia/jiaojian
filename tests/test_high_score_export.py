import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from process_data import (
    LONG_ORDER_SCORE_SCENE,
    PICKUP_SCORE_SCENE,
    attributed_deduction_fallback,
    build_deduction_customer_index,
    build_latest_customer_metadata,
    build_long_order_volume_calculator,
    write_high_score_excel,
)


class HighScoreExportTests(unittest.TestCase):
    def test_long_order_level_counts_zero_but_not_missing_or_outside_window(self):
        summarize = build_long_order_volume_calculator({
            "网点甲": {
                "2026-08-31": 5000,
                "2026-09-01": 0,
                "2026-09-02": None,
                "2026-09-15": 200,
            },
        })
        self.assertEqual(summarize("网点甲", "2026-09-15"), {
            "level": "100-500", "average": 100, "days": 2,
        })
        self.assertEqual(summarize("网点乙", "2026-09-15"), {
            "level": "—", "average": None, "days": 0,
        })

    def test_fallback_needs_one_customer_per_scene_and_known_status(self):
        index = {
            "网点甲": {
                "2026-09-15": {("code", "1"): {"是"}},
                "2026-09-14": {("code", "1"): {"是"}},
                "2026-09-13": {("code", "1"): {"是"}},
            },
        }
        self.assertEqual(attributed_deduction_fallback("网点甲", "2026-09-15", {PICKUP_SCORE_SCENE, LONG_ORDER_SCORE_SCENE}, index), "是")
        index["网点甲"]["2026-09-14"][("code", "2")] = {"否"}
        self.assertEqual(attributed_deduction_fallback("网点甲", "2026-09-15", {LONG_ORDER_SCORE_SCENE}, index), "-")
        index["网点甲"]["2026-09-14"].pop(("code", "2"))
        index["网点甲"]["2026-09-15"] = {("code", "2"): {"否"}}
        self.assertEqual(attributed_deduction_fallback("网点甲", "2026-09-15", {PICKUP_SCORE_SCENE, LONG_ORDER_SCORE_SCENE}, index), "-")
        index["网点甲"]["2026-09-15"] = {("code", "1"): {"未知"}}
        self.assertEqual(attributed_deduction_fallback("网点甲", "2026-09-15", {PICKUP_SCORE_SCENE}, index), "-")

    def test_fallback_uses_latest_known_status_for_customer_code(self):
        rows = [
            {
                "platform": "抖音", "branch": "网点甲", "customer": "客户甲", "customer_code": "C1",
                "date": "2026-09-10", "timeout_36h": 1, "has_shipping_fallback": "否",
            },
            {
                "platform": "抖音", "branch": "网点乙", "customer": "客户乙", "customer_code": "C1",
                "date": "2026-09-15", "timeout_36h": 1, "has_shipping_fallback": "是",
            },
        ]
        index = build_deduction_customer_index(rows, build_latest_customer_metadata(rows))
        self.assertEqual(index["网点甲"]["2026-09-10"][("code", "C1")], {"是"})

    def test_workbook_only_keeps_current_t1_t2_and_sorts_missing_after_zero(self):
        def row(branch, day, volume):
            return {
                "branch": branch,
                "stagnant_score": 6,
                "deduction_level": "100-",
                "deduction_average": 0,
                "deduction_days": 1,
                "long_order_level": "—",
                "long_order_average": None,
                "long_order_days": 0,
                "deduction_fallback": "-",
                "clearout_count": 0,
                "latest_score_date": day,
                "latest_deduction_volume": volume,
            }

        dashboard = {
            "platforms": {"抖音": {"dates": ["2026-09-14", "2026-09-15"]}},
            "high_scores_by_date": {
                "2026-09-14": [row("甲", "2026-09-13", 5)],
                "2026-09-15": [
                    row("乙", "2026-09-15", None),
                    row("丙", "2026-09-14", 0),
                    row("丁", "2026-09-13", 100),
                ],
            },
        }
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "export.xlsx"
            self.assertEqual(write_high_score_excel(path, dashboard), 2)
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = list(workbook.active.values)
                self.assertEqual([row[0].date().isoformat() for row in rows[1:]], ["2026-09-15", "2026-09-15"])
                self.assertEqual([row[2] for row in rows[1:]], ["丙", "乙"])
                self.assertEqual([row[1] for row in rows[1:]], ["T-2", "T-1"])
                self.assertEqual(rows[1][9], 0)
                self.assertIsNone(rows[2][9])
                self.assertEqual(rows[1][4], "100-\n日均 0 · 1天")
            finally:
                workbook.close()

    def test_workbook_level_text_matches_dashboard_display(self):
        dashboard = {
            "platforms": {"抖音": {"dates": ["2026-09-14", "2026-09-15"]}},
            "high_scores_by_date": {
                "2026-09-14": [],
                "2026-09-15": [{
                    "branch": "网点甲", "stagnant_score": 6,
                    "deduction_level": "1K-2K", "deduction_average": 1234.5, "deduction_days": 5,
                    "long_order_level": "100-500", "long_order_average": 286.49, "long_order_days": 3,
                    "deduction_fallback": "是", "clearout_count": 2,
                    "latest_score_date": "2026-09-15", "latest_deduction_volume": 10,
                }, {
                    "branch": "网点乙", "stagnant_score": 6,
                    "deduction_level": "超长单扣分", "deduction_average": None, "deduction_days": 1,
                    "long_order_level": "—", "long_order_average": None, "long_order_days": 0,
                    "deduction_fallback": "-", "clearout_count": 0,
                    "latest_score_date": "2026-09-14", "latest_deduction_volume": 0,
                }],
            },
        }
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "export.xlsx"
            self.assertEqual(write_high_score_excel(path, dashboard), 2)
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = list(workbook.active.values)
                self.assertEqual(rows[1][4], "1K-2K\n日均 1,235 · 5天")
                self.assertEqual(rows[1][5], "100-500\n日均 286 · 3天")
                self.assertNotIn("票", rows[1][4])
                self.assertEqual(rows[2][4], "超长单扣分")
                self.assertEqual(rows[2][5], "—")
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
