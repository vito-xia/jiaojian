import importlib.util
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from process_data import read_timeout


SCRIPT_PATH = Path(__file__).resolve().parents[1] / "数据源" / "脚本" / "T-1" / "更新T-1数据源.py"
spec = importlib.util.spec_from_file_location("t1_download", SCRIPT_PATH)
t1_download = importlib.util.module_from_spec(spec)
spec.loader.exec_module(t1_download)


def make_source(header_rows, invalid_group=False):
    workbook = Workbook()
    sheet = workbook.active
    for col, label in enumerate(("序号", "分部", "客户", "客户编码", "是否配置发运兜底", "是否历史有单无货"), 1):
        sheet.cell(1, col, label)
    if header_rows == 3:
        sheet.cell(1, 7, "网点交件")
        sheet.cell(1, 14, "网点直跑车")
        sheet.merge_cells("G1:M1")
        sheet.merge_cells("N1:P1")
        sheet.merge_cells("B1:B3")
        rate_label = "超时率"
    else:
        rate_label = "超时率B/C"
    groups = {
        7: "揽收->入首分拨(工单列)", 8: "揽收->入首分拨(超时列)",
        10: "揽收->入首分拨(第三列)", 11: "揽收->入首分拨(第四列)",
        12: "揽收->入首分拨(第五列)", 13: "揽收->入首分拨(第六列)",
    }
    for col, label in groups.items():
        sheet.cell(header_rows - 1, col, "未知指标" if invalid_group and col == 10 else label)
    sheet.cell(header_rows, 7, "票件量A")
    sheet.cell(header_rows, 8, "超时量B")
    sheet.cell(header_rows, 9, rate_label)
    for index in range(1, 201):
        row = header_rows + index
        sheet.cell(row, 2, "合计" if index == 1 else f"网点{index}")
        sheet.cell(row, 3, f"客户{index}")
        sheet.cell(row, 7, 0 if index == 2 else index)
        sheet.cell(row, 8, index)
        sheet.cell(row, 9, 0 if index == 2 else 1.5)
        sheet.cell(row, 10, index + 1)
        sheet.cell(row, 11, index + 2)
        sheet.cell(row, 12, index + 3)
        sheet.cell(row, 13, index + 4)
        if header_rows == 3:
            for col in (14, 15, 16):
                sheet.cell(row, col, 9999)
    payload = io.BytesIO()
    workbook.save(payload)
    workbook.close()
    return payload.getvalue()


class T1TemplateTests(unittest.TestCase):
    def test_both_templates_keep_same_body_rows_without_parsing_header(self):
        for header_rows in (2, 3):
            with self.subTest(header_rows=header_rows), tempfile.TemporaryDirectory() as temp_dir:
                payload = make_source(header_rows)
                with patch.object(t1_download, "OUT_DIR", Path(temp_dir)), patch.object(
                    t1_download.urllib.request, "urlopen", return_value=io.BytesIO(payload)
                ):
                    saved_rows = t1_download.download_and_save("抖音_9月23日.xlsx", "https://unused.example")
                self.assertEqual(saved_rows, header_rows + 198)

                records = read_timeout(Path(temp_dir), 2026)
                self.assertEqual(len(records), 197)
                self.assertEqual(records[0]["branch"], "网点2")
                self.assertEqual(records[0]["timeout_24h"], 0)
                self.assertEqual(records[0]["timeout_rate_36h"], 0)
                self.assertEqual(records[0]["timeout_120h"], 6)
                self.assertEqual(records[-1]["branch"], "网点198")
                self.assertEqual(records[-1]["timeout_48h"], 199)

    def test_unknown_metric_layout_fails_before_writing(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            with patch.object(t1_download, "OUT_DIR", Path(temp_dir)), patch.object(
                t1_download.urllib.request, "urlopen", return_value=io.BytesIO(make_source(3, invalid_group=True))
            ):
                with self.assertRaisesRegex(ValueError, "列顺序"):
                    t1_download.download_and_save("抖音_9月23日.xlsx", "https://unused.example")
            self.assertFalse(list(Path(temp_dir).iterdir()))


if __name__ == "__main__":
    unittest.main()
