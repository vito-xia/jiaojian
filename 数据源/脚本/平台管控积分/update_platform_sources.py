"""从两类系统导出更新④⑤⑥；--check 仅预演，正式更新先备份再替换。"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import sys
import tempfile
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel


PROJECT_ROOT = Path(__file__).resolve().parents[3]
DELIVERY_SCENE = "物流停滞-派送端"
TARGETS = {
    "④": ("④平台近一个月每日积分.xlsx", ("异常信息",)),
    "⑤": ("⑤平台预警&管控结果.xlsx", ("Sheet1",)),
    "⑥": ("⑥派送积分监控.xlsx", ("每日派送管控", "每日派送积分")),
}


class SourceUpdateError(RuntimeError):
    pass


def label(value: Any) -> str:
    return "" if value is None else str(value).strip()


def header_map(headers: tuple[str, ...], context: str) -> dict[str, int]:
    result = {}
    for index, name in enumerate(headers):
        key = name.casefold()
        if not key or key in result:
            raise SourceUpdateError(f"{context}：表头为空或重复（第 {index + 1} 列）")
        result[key] = index
    return result


def latest_export(directory: Path) -> Path:
    if not directory.is_dir():
        raise SourceUpdateError(f"缺少导出文件夹：{directory.name}")
    files = [p for p in directory.iterdir() if p.is_file()
             and p.suffix.lower() == ".xlsx" and not p.name.startswith("~$")]
    if not files:
        raise SourceUpdateError(f"{directory.name} 中没有可用的 .xlsx 导出，请补齐后重试")
    return max(files, key=lambda p: (p.stat().st_mtime_ns, p.name))


def signature(path: Path) -> dict[str, Any]:
    before = path.stat()
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    after = path.stat()
    if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
        raise SourceUpdateError(f"读取期间文件变化：{path.name}")
    return {"size": after.st_size, "mtime_ns": after.st_mtime_ns, "sha256": digest}


def parse_day(value: Any, epoch: datetime, context: str) -> date:
    try:
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            converted = from_excel(value, epoch)
            if isinstance(converted, datetime):
                return converted.date()
        return datetime.fromisoformat(label(value).replace("/", "-")).date()
    except (TypeError, ValueError, OverflowError):
        pass
    raise SourceUpdateError(f"{context}：无法识别日期 {value!r}")


@dataclass
class Export:
    path: Path
    sheet: str
    headers: tuple[str, ...]
    rows: list[tuple[int, tuple[Any, ...]]]
    epoch: datetime


def read_export(path: Path, required: set[str]) -> Export:
    with path.open("rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=False)
        try:
            candidates = []
            for sheet in workbook:
                # 部分系统导出把 dimension 错写为 A1:A1，必须忽略该标记。
                sheet.reset_dimensions()
                first = next(sheet.iter_rows(), ())
                headers = tuple(label(c.value) for c in first)
                if required <= {h.casefold() for h in headers}:
                    candidates.append((sheet, headers))
            if len(candidates) != 1:
                raise SourceUpdateError(f"{path.name}：应有且仅有一个包含必要表头的工作表，实际 {len(candidates)} 个")
            sheet, headers = candidates[0]
            header_map(headers, path.name)
            rows = []
            for row_index, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(headers)), 2):
                values = tuple(c.value.strip() if isinstance(c.value, str) else c.value for c in cells)
                if all(v in (None, "") for v in values):
                    continue
                if any(c.data_type == "f" for c in cells):
                    raise SourceUpdateError(f"{path.name} 第 {row_index} 行含公式，应使用系统导出的原始数值")
                rows.append((row_index, values))
            return Export(path, sheet.title, headers, rows, workbook.epoch)
        finally:
            workbook.close()


@dataclass
class Record:
    values: tuple[Any, ...]
    day: date
    styles: tuple[Any, ...]
    dimension: Any = None
    comments: dict[int, Any] = field(default_factory=dict)
    hyperlinks: dict[int, Any] = field(default_factory=dict)


@dataclass
class SheetUpdate:
    path: Path
    sheet: str
    headers: tuple[str, ...]
    rows: list[Record]
    date_formats: dict[int, str]
    summary: dict[str, Any]


@dataclass
class UpdatePlan:
    manual_dir: Path
    workbooks: dict[Path, Any]
    updates: list[SheetUpdate]
    inputs: dict[Path, dict[str, Any]]
    exports: dict[str, Path]

    def close(self) -> None:
        for workbook in self.workbooks.values():
            workbook.close()


def old_records(sheet: Any, epoch: datetime, date_name: str, context: str) -> tuple[tuple[str, ...], list[Record]]:
    headers = tuple(label(c.value) for c in sheet[1])
    columns = header_map(headers, context)
    if date_name.casefold() not in columns:
        raise SourceUpdateError(f"{context}：缺少 {date_name} 列")
    records = []
    for index, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(headers)), 2):
        values = tuple(c.value for c in cells)
        if all(v in (None, "") for v in values):
            continue
        day = parse_day(values[columns[date_name.casefold()]], epoch, f"{context} 第 {index} 行")
        records.append(Record(
            values, day, tuple(copy(c._style) for c in cells), copy(sheet.row_dimensions.get(index)),
            {i: copy(c.comment) for i, c in enumerate(cells) if c.comment},
            {i: copy(c.hyperlink) for i, c in enumerate(cells) if c.hyperlink},
        ))
    return headers, records


def make_update(path: Path, workbook: Any, sheet_name: str, source: Export,
                delivery: bool, score_dates: set[date] | None = None) -> SheetUpdate:
    sheet = workbook[sheet_name]
    is_score = score_dates is not None
    date_name = "违规日期" if is_score else "开始日期"
    headers, old = old_records(sheet, workbook.epoch, date_name, f"{path.name}/{sheet_name}")
    target_columns = header_map(headers, sheet_name)
    source_columns = header_map(source.headers, source.path.name)
    missing = [h for h in headers if h.casefold() not in source_columns]
    if missing:
        raise SourceUpdateError(f"{source.path.name}：缺少目标列 {', '.join(missing)}")
    date_column = target_columns[date_name.casefold()]
    styles = old[0].styles if old else tuple(copy(sheet.cell(2, i + 1)._style) for i in range(len(headers)))
    dimension = old[0].dimension if old else None
    date_formats = {}
    if not is_score:
        for name in ("开始日期", "预计完结日期"):
            col = target_columns[name]
            date_formats[col] = next((sheet.cell(i, col + 1).number_format
                                     for i in range(2, sheet.max_row + 1)
                                     if isinstance(sheet.cell(i, col + 1).value, (date, datetime))), "mm-dd-yy")
    incoming = []
    for row_index, row in source.rows:
        if is_score and label(row[source_columns["涉及商家名称"]]) != "全部商家":
            continue
        scene = label(row[source_columns["违规场景"]])
        if (DELIVERY_SCENE in scene) != delivery:
            continue
        values = [row[source_columns[h.casefold()]] for h in headers]
        day = parse_day(values[date_column], source.epoch, f"{source.path.name} 第 {row_index} 行")
        if is_score:
            values[date_column] = day.isoformat()
        else:
            values[date_column] = datetime.combine(day, time())
            end_column = target_columns["预计完结日期"]
            if label(values[end_column]) not in ("", "-", "--", "—", "/"):
                end = parse_day(values[end_column], source.epoch, f"{source.path.name} 第 {row_index} 行预计完结日期")
                values[end_column] = datetime.combine(end, time())
        incoming.append(Record(tuple(values), day, styles, dimension))
    cutoff = None if is_score or not incoming else min(r.day for r in incoming)
    if is_score:
        # 整份筛选后的日期集合同时应用于④和⑥，某场景为零条也要清除旧日数据。
        if old:
            last_day = max(r.day for r in old)
            if last_day not in score_dates:
                raise SourceUpdateError(f"{path.name}/{sheet_name}：积分导出未包含旧表最后一天 {last_day}，请重新导出完整回补日期")
            if max(score_dates) < last_day:
                raise SourceUpdateError(f"{source.path.name}：积分导出早于旧表，停止更新")
        retained = [r for r in old if r.day not in score_dates]
    else:
        retained = [r for r in old if cutoff is None or r.day < cutoff]
    changed_statuses = 0
    if not is_score:
        status_column = target_columns["管控状态"]
        for record in retained:
            values = list(record.values)
            changed_statuses += values[status_column] != "已完结"
            values[status_column] = "已完结"
            record.values = tuple(values)
    rows = sorted(retained + incoming, key=lambda record: record.day)
    old_dates = {r.day for r in old}
    summary = {
        "file": path.name, "sheet": sheet_name, "kind": "积分" if is_score else "管控",
        "old_rows": len(old), "deleted_rows": len(old) - len(retained),
        "retained_rows": len(retained), "completed_status_rows": changed_statuses,
        "incoming_rows": len(incoming), "final_rows": len(rows),
        "cutoff": cutoff.isoformat() if cutoff else None,
        "replace_dates": sorted(d.isoformat() for d in score_dates) if is_score else [],
        "new_dates": sorted({r.day.isoformat() for r in incoming if r.day not in old_dates}),
    }
    return SheetUpdate(path, sheet_name, headers, rows, date_formats, summary)


def plan_sources(data_dir: Path) -> UpdatePlan:
    manual = data_dir / "数据源-手动更新"
    exports = {"管控": latest_export(manual / "系统导出-管控"), "积分": latest_export(manual / "系统管控-积分")}
    paths = [manual / name for name, _ in TARGETS.values()]
    for path in paths:
        if not path.is_file():
            raise SourceUpdateError(f"缺少历史目标表：{path.name}")
    inputs = {p: signature(p) for p in [*exports.values(), *paths]}
    controls = read_export(exports["管控"], {"管控单id", "违规场景", "开始日期", "管控状态", "预计完结日期"})
    scores = read_export(exports["积分"], {"网点名称", "涉及商家名称", "违规场景", "违规日期", "当前违规积分"})
    if not controls.rows:
        raise SourceUpdateError("管控导出没有正文，无法确认有效快照，停止更新")
    control_columns = header_map(controls.headers, controls.path.name)
    seen = set()
    for index, row in controls.rows:
        control_id = label(row[control_columns["管控单id"]])
        if not control_id or control_id in seen:
            raise SourceUpdateError(f"{controls.path.name} 第 {index} 行管控单id为空或重复")
        seen.add(control_id)
    score_columns = header_map(scores.headers, scores.path.name)
    dates = {parse_day(row[score_columns["违规日期"]], scores.epoch, f"{scores.path.name} 第 {index} 行")
             for index, row in scores.rows if label(row[score_columns["涉及商家名称"]]) == "全部商家"}
    if not dates:
        raise SourceUpdateError("积分导出中没有“全部商家”的有效日期，停止更新")
    workbooks = {}
    try:
        for path, (_, required_sheets) in zip(paths, TARGETS.values()):
            workbook = load_workbook(path, data_only=False)
            workbooks[path] = workbook
            if not set(required_sheets) <= set(workbook.sheetnames):
                raise SourceUpdateError(f"{path.name}：缺少工作表 {', '.join(required_sheets)}")
        updates = [
            make_update(paths[0], workbooks[paths[0]], "异常信息", scores, False, dates),
            make_update(paths[1], workbooks[paths[1]], "Sheet1", controls, False),
            make_update(paths[2], workbooks[paths[2]], "每日派送管控", controls, True),
            make_update(paths[2], workbooks[paths[2]], "每日派送积分", scores, True, dates),
        ]
        return UpdatePlan(manual, workbooks, updates, inputs, exports)
    except Exception:
        for workbook in workbooks.values():
            workbook.close()
        raise


def apply_sheet(workbook: Any, update: SheetUpdate) -> None:
    sheet = workbook[update.sheet]
    sheet.delete_rows(2, max(0, sheet.max_row - 1))
    for index in list(sheet.row_dimensions):
        if index >= 2:
            del sheet.row_dimensions[index]
    for index, record in enumerate(update.rows, 2):
        for col, value in enumerate(record.values):
            cell = sheet.cell(index, col + 1, value)
            cell._style = copy(record.styles[col])
            if col in update.date_formats:
                cell.number_format = update.date_formats[col]
            if col in record.comments:
                cell.comment = copy(record.comments[col])
            if col in record.hyperlinks:
                cell.hyperlink = copy(record.hyperlinks[col])
        if record.dimension is not None:
            dim = copy(record.dimension)
            dim.index = index
            sheet.row_dimensions[index] = dim
    if sheet.auto_filter.ref:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(update.headers))}{len(update.rows) + 1}"


def comparable(value: Any) -> Any:
    return None if value == "" else value


def validate_staged(path: Path, updates: list[SheetUpdate], expected_sheets: list[str]) -> None:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != expected_sheets:
            raise SourceUpdateError("生成后的工作表名称或顺序发生变化")
        for update in updates:
            sheet = workbook[update.sheet]
            actual = sheet.iter_rows(values_only=True)
            if tuple(next(actual)) != update.headers:
                raise SourceUpdateError(f"{update.sheet}：生成后表头不一致")
            for index, expected in enumerate(update.rows, 2):
                row = next(actual, None)
                if row is None or tuple(map(comparable, row)) != tuple(map(comparable, expected.values)):
                    raise SourceUpdateError(f"{update.sheet} 第 {index} 行：保存后数值不一致")
            if next(actual, None) is not None:
                raise SourceUpdateError(f"{update.sheet}：生成后存在多余行")
    finally:
        workbook.close()


def check_inputs(plan: UpdatePlan) -> None:
    for kind, path in plan.exports.items():
        if latest_export(path.parent) != path:
            raise SourceUpdateError(f"处理期间新增了更晚的{kind}导出，请重新运行")
    for path, before in plan.inputs.items():
        if signature(path) != before:
            raise SourceUpdateError(f"处理期间文件发生变化：{path.name}，停止替换")


def write_report(path: Path, report: dict[str, Any]) -> None:
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def commit_plan(plan: UpdatePlan, backup_root: Path) -> Path:
    lock_path = plan.manual_dir / ".platform_sources_update.lock"
    try:
        lock_fd = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError:
        raise SourceUpdateError("已有积分管控更新正在运行；若上次异常退出，请先核对备份，再移除更新锁文件") from None
    staged = {}
    replaced = []
    backup_dir = None
    try:
        os.close(lock_fd)
        check_inputs(plan)
        for update in plan.updates:
            apply_sheet(plan.workbooks[update.path], update)
        for path, workbook in plan.workbooks.items():
            fd, name = tempfile.mkstemp(prefix=".platform-update-", suffix=".xlsx", dir=plan.manual_dir)
            os.close(fd)
            staged[path] = Path(name)
            workbook.save(name)
            validate_staged(staged[path], [u for u in plan.updates if u.path == path], workbook.sheetnames)
        check_inputs(plan)
        backup_root.mkdir(parents=True, exist_ok=True)
        backup_dir = backup_root / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        backup_dir.mkdir()
        for path in plan.workbooks:
            shutil.copy2(path, backup_dir / path.name)
            if signature(backup_dir / path.name)["sha256"] != plan.inputs[path]["sha256"]:
                raise SourceUpdateError(f"备份校验失败：{path.name}")
        report = {
            "status": "prepared", "generated_at": datetime.now().isoformat(timespec="seconds"),
            "exports": {k: {"file": str(p), **plan.inputs[p]} for k, p in plan.exports.items()},
            "targets": [{"file": str(p), **plan.inputs[p]} for p in plan.workbooks],
            "updates": [u.summary for u in plan.updates],
        }
        report_path = backup_dir / "更新报告.json"
        write_report(report_path, report)
        check_inputs(plan)
        try:
            for path, temporary in staged.items():
                os.replace(temporary, path)
                replaced.append(path)
            report["status"] = "completed"
            write_report(report_path, report)
        except Exception as error:
            recovery_errors = []
            for path in reversed(replaced):
                try:
                    shutil.copy2(backup_dir / path.name, path)
                except Exception as recovery_error:
                    recovery_errors.append(f"{path.name}: {recovery_error}")
            report["status"] = "recovery_required" if recovery_errors else "rolled_back"
            report["error"] = str(error)
            report["recovery_errors"] = recovery_errors
            try:
                write_report(report_path, report)
            except OSError:
                pass
            message = "恢复失败，请使用备份手动恢复" if recovery_errors else "已恢复本次替换前的文件"
            raise SourceUpdateError(f"替换失败，{message}；备份：{backup_dir}；原因：{error}") from error
        return backup_dir
    finally:
        for path in staged.values():
            path.unlink(missing_ok=True)
        lock_path.unlink(missing_ok=True)


def print_plan(plan: UpdatePlan) -> None:
    for kind, path in plan.exports.items():
        print(f"{kind}导出：{path.name}")
    for update in plan.updates:
        s = update.summary
        dates = ", ".join(s["replace_dates"]) if s["kind"] == "积分" else f">= {s['cutoff']}" if s["cutoff"] else "无新记录，保留旧记录并完结"
        print(f"{s['file']} / {s['sheet']}：覆盖 {dates}；旧 {s['old_rows']}，删除 {s['deleted_rows']}，"
              f"导入 {s['incoming_rows']}，状态改完结 {s['completed_status_rows']}，最终 {s['final_rows']}")
        if s["new_dates"]:
            print(f"  新增日期：{', '.join(s['new_dates'])}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", type=Path, default=PROJECT_ROOT / "数据源")
    parser.add_argument("--backup-dir", type=Path, help="默认：数据源/脚本处理后输出/平台管控积分-备份")
    parser.add_argument("--check", action="store_true", help="仅校验并预演，不写表、不创建备份")
    args = parser.parse_args()
    plan = None
    try:
        plan = plan_sources(args.data_dir.resolve())
        print_plan(plan)
        if args.check:
            check_inputs(plan)
            print("预演通过，未写入任何文件。")
        else:
            backup_root = args.backup_dir or args.data_dir / "脚本处理后输出" / "平台管控积分-备份"
            backup = commit_plan(plan, backup_root)
            print(f"④⑤⑥更新完成；备份和更新报告：{backup}")
        return 0
    except Exception as error:
        print(f"积分管控更新失败：{error}", file=sys.stderr)
        return 1
    finally:
        if plan is not None:
            plan.close()


if __name__ == "__main__":
    raise SystemExit(main())
