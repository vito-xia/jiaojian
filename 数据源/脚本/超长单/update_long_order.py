"""截取超长单列表并生成抖音网点趋势分片。

这是暂时独立于 ``process_data.py`` 的验证脚本。它读取
``数据源/数据源-手动更新/超长单-手动更新`` 中每个日期工作簿的 ``列表数据``
工作表，先按 ``超长单异常运单数`` 降序，再保留 TOP1000 数据，精简副本写入历史处理后目录，同时生成供页面按需加载的
``data/dashboard_long_order.js``。分片中的每个趋势点保留超长单监控表所需的机构、
业务省区、应签总数和异常等级字段；源表中的省份、城市只保留在处理后 Excel 中。

源工作簿的 XML 维度可能错误地写成 A1:A1，因此读取前必须调用
``reset_dimensions()``，不能依据 ``max_row`` 判断数据行数。
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import tempfile
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


TARGET_SHEET = "列表数据"
PROVINCE_HEADER = "省份"
CITY_HEADER = "城市"
BRANCH_HEADER = "网点"
COUNT_HEADER = "超长单异常运单数"
EXPECTED_SIGN_COUNT_HEADER = "超长单应签运单总数"
RATE_HEADER = "超长单异常率"
ABNORMAL_LEVEL_HEADER = "超长单异常率-异常等级"
DATA_ROW_LIMIT = 1000
CHUNK_NAME = "long-order"
PAYLOAD_SCHEMA_VERSION = 3
SOURCE_FILE_PATTERN = re.compile(r"^(?P<month>\d{1,2})月(?P<day>\d{1,2})日\.xlsx$", re.IGNORECASE)
EMPTY_VALUES = {"", "-", "--", "—", "/", "无", "暂无", "null", "none", "nan"}


class LongOrderFormatError(RuntimeError):
    """源表存在但无法安全转换。"""


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_number(value: Any) -> float | None:
    raw = text(value)
    if raw.lower() in EMPTY_VALUES:
        return None
    cleaned = raw.replace(",", "").replace("，", "").replace("%", "")
    try:
        return float(cleaned)
    except (TypeError, ValueError):
        return None


def normalize_count(value: Any) -> int | float | None:
    amount = parse_number(value)
    if amount is None:
        return None
    return int(amount) if amount.is_integer() else round(amount, 4)


def normalize_rate_points(value: Any) -> float | None:
    """将源表比例统一成页面使用的百分比点数。"""
    raw = text(value)
    amount = parse_number(value)
    if amount is None:
        return None
    if "%" not in raw and abs(amount) <= 1:
        amount *= 100
    return round(amount, 4)


def project_root_from_script() -> Path:
    # .../数据源/超长单/脚本/update_long_order.py
    return Path(__file__).resolve().parents[3]


def default_year(project_root: Path) -> int:
    report_path = project_root / "data" / "data_quality_report.json"
    try:
        report = json.loads(report_path.read_text(encoding="utf-8"))
        as_of = text(report.get("as_of"))
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", as_of):
            return int(as_of[:4])
    except (OSError, ValueError, TypeError):
        pass
    return date.today().year


def load_business_province_mapping(path: Path) -> dict[str, str]:
    """读取 process_data.py 生成的网点到业务省区映射。"""
    if not path.is_file():
        raise FileNotFoundError(f"未找到网点归属映射，请先运行 process_data.py：{path}")
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as error:
        raise LongOrderFormatError(f"网点归属映射读取失败：{path}") from error
    if not isinstance(payload, dict):
        raise LongOrderFormatError(f"网点归属映射格式错误，根节点必须是对象：{path}")
    mapping: dict[str, str] = {}
    for branch, item in payload.items():
        if not isinstance(item, dict):
            continue
        branch_name = text(branch)
        if branch_name:
            mapping[branch_name] = text(item.get("province"))
    return mapping


def parse_source_day(path: Path, year: int) -> str:
    match = SOURCE_FILE_PATTERN.fullmatch(path.name)
    if not match:
        raise LongOrderFormatError(f"文件名必须是 M月D日.xlsx：{path.name}")
    try:
        return date(year, int(match.group("month")), int(match.group("day"))).isoformat()
    except ValueError as error:
        raise LongOrderFormatError(f"文件名日期无效：{path.name}") from error


def source_files(source_dir: Path, year: int) -> list[tuple[str, Path]]:
    if not source_dir.exists():
        raise FileNotFoundError(f"未找到源目录：{source_dir}")
    pairs: list[tuple[str, Path]] = []
    for path in source_dir.iterdir():
        if not path.is_file() or path.suffix.lower() != ".xlsx" or path.name.startswith("~$"):
            continue
        day = parse_source_day(path, year)
        lock_path = path.with_name("~$" + path.name)
        if lock_path.exists() and is_file_locked(path):
            raise LongOrderFormatError(f"源文件正在被 Excel 锁定，请先关闭：{path.name}")
        if lock_path.exists():
            print(f"提示：发现陈旧锁文件，正式文件当前可读，已忽略：{lock_path.name}")
        pairs.append((day, path))
    if not pairs:
        raise FileNotFoundError(f"源目录中没有可处理的 M月D日.xlsx：{source_dir}")
    pairs.sort(key=lambda item: item[0])
    seen: set[str] = set()
    for day, path in pairs:
        if day in seen:
            raise LongOrderFormatError(f"同一日期存在多个源文件：{day}")
        seen.add(day)
    return pairs


def is_file_locked(path: Path) -> bool:
    """用独占读测试确认锁文件是否对应真实占用，避免误伤陈旧 ~$ 文件。"""
    if os.name == "nt":
        import ctypes
        from ctypes import wintypes

        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        create_file = kernel32.CreateFileW
        create_file.argtypes = [
            wintypes.LPCWSTR,
            wintypes.DWORD,
            wintypes.DWORD,
            wintypes.LPVOID,
            wintypes.DWORD,
            wintypes.DWORD,
            wintypes.HANDLE,
        ]
        create_file.restype = wintypes.HANDLE
        close_handle = kernel32.CloseHandle
        close_handle.argtypes = [wintypes.HANDLE]
        close_handle.restype = wintypes.BOOL
        handle = create_file(
            str(path),
            0x80000000,  # GENERIC_READ
            0,  # no sharing: fail while Excel holds the file
            None,
            3,  # OPEN_EXISTING
            0x80,  # FILE_ATTRIBUTE_NORMAL
            None,
        )
        invalid_handle = wintypes.HANDLE(-1).value
        if handle == invalid_handle:
            return True
        close_handle(handle)
        return False
    try:
        descriptor = os.open(path, os.O_RDONLY)
        os.close(descriptor)
    except OSError:
        return True
    return False


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def header_info(rows: list[tuple[Any, ...]], path: Path) -> tuple[int, list[str], dict[str, int]]:
    for row_number, row in enumerate(rows, 1):
        headers = [text(value) for value in row]
        while headers and not headers[-1]:
            headers.pop()
        required = {
            PROVINCE_HEADER,
            CITY_HEADER,
            BRANCH_HEADER,
            COUNT_HEADER,
            EXPECTED_SIGN_COUNT_HEADER,
            RATE_HEADER,
            ABNORMAL_LEVEL_HEADER,
        }
        if required.issubset(headers):
            duplicates = [header for header in required if headers.count(header) != 1]
            if duplicates:
                raise LongOrderFormatError(f"{path.name} 表头重复：{', '.join(duplicates)}")
            columns = {header: headers.index(header) for header in required}
            return row_number, headers, columns
    raise LongOrderFormatError(
        f"{path.name} 的 {TARGET_SHEET} 前30行未找到表头："
        f"{PROVINCE_HEADER}、{CITY_HEADER}、{BRANCH_HEADER}、{COUNT_HEADER}、"
        f"{EXPECTED_SIGN_COUNT_HEADER}、{RATE_HEADER}、{ABNORMAL_LEVEL_HEADER}"
    )


def read_one_source(day: str, path: Path, row_limit: int) -> dict[str, Any]:
    before_hash = sha256_file(path)
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if TARGET_SHEET not in workbook.sheetnames:
            raise LongOrderFormatError(f"{path.name} 缺少工作表：{TARGET_SHEET}")
        sheet = workbook[TARGET_SHEET]
        # 部分下载工作簿将 dimension 错误写成 A1:A1；这里是必需的修复动作。
        sheet.reset_dimensions()
        first_rows = list(sheet.iter_rows(min_row=1, max_row=30, values_only=True))
        header_row, headers, columns = header_info(first_rows, path)
        source_rows: list[tuple[int, list[Any], int | float | None]] = []
        branch_rows: dict[str, dict[str, Any]] = {}
        warnings: list[str] = []
        for row_number, row in enumerate(
            sheet.iter_rows(
                min_row=header_row + 1,
                max_col=len(headers),
                values_only=True,
            ),
            header_row + 1,
        ):
            values = list(row[: len(headers)])
            if len(values) < len(headers):
                values.extend([None] * (len(headers) - len(values)))
            if not any(value not in (None, "") for value in values):
                continue
            count = normalize_count(values[columns[COUNT_HEADER]])
            source_rows.append((row_number, values, count))

        # 先完成全表排序，再取 TOP1000。None 放在真实 0 后面，排序稳定以保留源表顺序。
        source_rows.sort(
            key=lambda item: (
                item[2] is not None,
                item[2] if item[2] is not None else 0,
            ),
            reverse=True,
        )
        selected_rows = source_rows[:row_limit]
        output_rows = []
        for source_rank, (row_number, values, count) in enumerate(selected_rows, 1):
            output_rows.append(values)
            branch = text(values[columns[BRANCH_HEADER]])
            if not branch:
                continue
            if branch in branch_rows:
                raise LongOrderFormatError(f"{path.name} 排序后 TOP{row_limit} 内网点重复：{branch}")
            expected_sign_count = normalize_count(values[columns[EXPECTED_SIGN_COUNT_HEADER]])
            rate = normalize_rate_points(values[columns[RATE_HEADER]])
            if count is None or expected_sign_count is None or rate is None:
                warnings.append(f"第{row_number}行量、应签总数或率为空")
            branch_rows[branch] = {
                "date": day,
                "source_rank": source_rank,
                "branch": branch,
                "abnormal_count": count,
                "expected_sign_count": expected_sign_count,
                "abnormal_rate": rate,
                "abnormal_level": text(values[columns[ABNORMAL_LEVEL_HEADER]]) or None,
            }
        if not output_rows:
            raise LongOrderFormatError(f"{path.name} 没有可保留的数据行")
    finally:
        workbook.close()
    after_hash = sha256_file(path)
    if before_hash != after_hash:
        raise LongOrderFormatError(f"读取期间源文件发生变化：{path.name}")
    return {
        "date": day,
        "source_path": path,
        "source_sha256": before_hash,
        "headers": headers,
        "header_row": header_row,
        "rows": output_rows,
        "branch_rows": branch_rows,
        "warnings": warnings,
    }


def write_trimmed_workbook(path: Path, headers: list[str], rows: list[list[Any]]) -> None:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(TARGET_SHEET)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def build_payload(
    records: list[dict[str, Any]],
    generated_at: str,
    year: int,
    row_limit: int,
    business_provinces: dict[str, str],
) -> dict[str, Any]:
    trends: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        for branch, point in record["branch_rows"].items():
            enriched = dict(point)
            enriched["business_province"] = business_provinces.get(branch, "")
            trends[branch].append(enriched)
    for points in trends.values():
        points.sort(key=lambda item: item["date"])
    ordered_trends = {branch: trends[branch] for branch in sorted(trends)}
    return {
        "long_order_trends": ordered_trends,
        "long_order_meta": {
            "schema_version": PAYLOAD_SCHEMA_VERSION,
            "generated_at": generated_at,
            "year": year,
            "row_limit": row_limit,
            "source_file_count": len(records),
            "source_dates": [record["date"] for record in records],
            "source_row_count": sum(len(record["rows"]) for record in records),
            "source_branch_row_count": sum(len(record["branch_rows"]) for record in records),
            "warning_count": sum(len(record["warnings"]) for record in records),
        },
    }


def write_js_payload(path: Path, payload: dict[str, Any]) -> None:
    encoded = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    content = (
        f"window.__JIAOJIAN_REGISTER_CHUNK__({json.dumps(CHUNK_NAME, ensure_ascii=False)},{encoded});\n"
    )
    path.write_text(content, encoding="utf-8", newline="\n")


def write_manifest(path: Path, records: list[dict[str, Any]], generated_at: str, year: int, row_limit: int) -> None:
    manifest = {
        "schema_version": 1,
        "generated_at": generated_at,
        "year": year,
        "row_limit": row_limit,
        "entries": [
            {
                "date": record["date"],
                "source": record["source_path"].name,
                "source_sha256": record["source_sha256"],
                "header_row": record["header_row"],
                "output_rows": len(record["rows"]),
                "branch_rows": len(record["branch_rows"]),
                "warnings": record["warnings"],
            }
            for record in records
        ],
    }
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")


def collect_records(source_dir: Path, processed_dir: Path, year: int, row_limit: int = DATA_ROW_LIMIT) -> list[dict[str, Any]]:
    """Read current files plus retained history, preferring the current file for duplicate dates."""
    pairs = source_files(source_dir, year)
    current_records = [read_one_source(day, path, row_limit) for day, path in pairs]
    current_dates = {record["date"] for record in current_records}
    history_records: list[dict[str, Any]] = []
    if processed_dir.is_dir() and any(
        path.is_file() and path.suffix.lower() == ".xlsx" and not path.name.startswith("~$")
        for path in processed_dir.iterdir()
    ):
        history_pairs = [
            (day, path)
            for day, path in source_files(processed_dir, year)
            if day not in current_dates
        ]
        history_records = [read_one_source(day, path, row_limit) for day, path in history_pairs]
    return sorted(history_records + current_records, key=lambda record: record["date"])


def run(args: argparse.Namespace) -> int:
    project_root = project_root_from_script()
    data_source_dir = project_root / "数据源"
    source_dir = args.source_dir or data_source_dir / "数据源-手动更新" / "超长单-手动更新"
    processed_dir = args.processed_dir or data_source_dir / "脚本处理后输出" / "超长单-脚本处理后"
    output_path = args.output or project_root / "data" / "dashboard_long_order.js"
    mapping_path = project_root / "data" / "branch_mapping.json"
    year = args.year or default_year(project_root)
    row_limit = DATA_ROW_LIMIT
    records = collect_records(source_dir, processed_dir, year, row_limit)
    business_provinces = load_business_province_mapping(mapping_path)
    current_records = [record for record in records if record["source_path"].parent == source_dir]
    generated_at = datetime.now().replace(microsecond=0).isoformat(sep=" ")
    payload = build_payload(records, generated_at, year, row_limit, business_provinces)
    summary = payload["long_order_meta"]
    print(
        f"源文件 {summary['source_file_count']} 个，日期 {summary['source_dates'][0]} 至 "
        f"{summary['source_dates'][-1]}，列表数据行 {summary['source_row_count']}，"
        f"网点日记录 {summary['source_branch_row_count']}。"
    )
    for record in records:
        if record["warnings"]:
            print(f"提示：{record['source_path'].name} 有 {len(record['warnings'])} 条量率缺失记录，保留为空值。")
    if args.check:
        print("校验完成，未写入处理后文件或页面分片。")
        return 0

    processed_dir.mkdir(parents=True, exist_ok=True)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    stage_parent = processed_dir.parent
    stage_parent.mkdir(parents=True, exist_ok=True)
    stage_dir = Path(tempfile.mkdtemp(prefix=".long-order-stage-", dir=str(stage_parent)))
    try:
        for record in current_records:
            write_trimmed_workbook(stage_dir / record["source_path"].name, record["headers"], record["rows"])
        write_manifest(stage_dir / "manifest.json", records, generated_at, year, row_limit)
        write_js_payload(stage_dir / output_path.name, payload)
        for record in current_records:
            target = processed_dir / record["source_path"].name
            os.replace(stage_dir / record["source_path"].name, target)
        os.replace(stage_dir / "manifest.json", processed_dir / "manifest.json")
        os.replace(stage_dir / output_path.name, output_path)
    finally:
        shutil.rmtree(stage_dir, ignore_errors=True)
    print(f"处理后文件已生成：{processed_dir}")
    print(f"页面分片已生成：{output_path}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="按超长单异常运单数降序截取列表 TOP1000 并生成抖音趋势分片")
    parser.add_argument("--year", type=int, help="文件名月份日期对应的年份，默认取当前数据 as_of 年份")
    parser.add_argument("--source-dir", type=Path, help="明细源目录")
    parser.add_argument("--processed-dir", type=Path, help="处理后目录")
    parser.add_argument("--output", type=Path, help="dashboard_long_order.js 输出路径")
    parser.add_argument("--check", action="store_true", help="只校验和汇总，不写入任何输出")
    parser.add_argument("--no-pause", action="store_true", help=argparse.SUPPRESS)
    args = parser.parse_args()
    try:
        return run(args)
    except Exception as error:
        print(f"超长单处理失败：{error}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
