# -*- coding: utf-8 -*-
"""
交件超时数据源更新工具
双击运行：读取「数据源-手动更新/交件链接清单T-1日.txt」中的下载链接，
逐个下载并保留表头及前198行正文，保存到 T-1 处理后目录。
兼容旧版两行表头和新版三行表头，保留原始列及格式。
看板刷新由可见的 BAT 入口统一编排，避免 Python 和 BAT 重复执行。
"""
import io
import urllib.request
from pathlib import Path
import openpyxl
import warnings

warnings.filterwarnings("ignore")

PROJECT_ROOT = Path(__file__).resolve().parents[3]
DATA_SOURCE_DIR = PROJECT_ROOT / "数据源"
OUT_DIR = DATA_SOURCE_DIR / "脚本处理后输出" / "交件T-1脚本处理后"
LINK_FILE = DATA_SOURCE_DIR / "数据源-手动更新" / "交件链接清单T-1日.txt"
KEEP_DATA_ROWS = 198  # 与旧版“前200行（两行表头）”的正文行数一致


def read_links(file_path):
    """解析链接清单：每行「文件名.xlsx 链接」或「平台 日期 链接」"""
    links = []
    with open(file_path, "r", encoding="utf-8-sig") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            parts = line.split()
            if len(parts) < 2:
                print(f"  跳过无法解析的行: {line}")
                continue
            url = parts[-1]
            name_part = " ".join(parts[:-1])
            if name_part.lower().endswith(".xlsx"):
                fname = name_part
            elif len(parts) == 3:
                # 平台 日期 链接
                fname = f"{parts[0]}_{parts[1]}.xlsx"
            else:
                print(f"  跳过无法解析的行: {line}")
                continue
            links.append((fname, url))
    return links


def copy_cell_style(src, dst):
    if src.has_style:
        dst.font = openpyxl.styles.Font(
            name=src.font.name, size=src.font.size, bold=src.font.bold,
            color=src.font.color)
        dst.alignment = openpyxl.styles.Alignment(
            horizontal=src.alignment.horizontal,
            vertical=src.alignment.vertical,
            wrap_text=src.alignment.wrap_text)
        dst.fill = openpyxl.styles.PatternFill(
            fill_type=src.fill.fill_type, fgColor=src.fill.fgColor)
        dst.border = openpyxl.styles.Border(
            left=src.border.left, right=src.border.right,
            top=src.border.top, bottom=src.border.bottom)
        dst.number_format = src.number_format


def header_row_count(ws):
    """识别 T-1 两版模板，避免把新版第 3 行表头当成业务数据。"""
    basic_headers = ("序号", "分部", "客户", "客户编码", "是否配置发运兜底", "是否历史有单无货")
    if tuple(ws.cell(1, col).value for col in range(1, 7)) != basic_headers:
        raise ValueError("T-1 文件基础表头不符合已知模板")

    metric_headers = ("票件量A", "超时量B")
    if tuple(ws.cell(2, col).value for col in (7, 8)) == metric_headers:
        header_rows = 2
        expected_rate = "超时率B/C"
    elif tuple(ws.cell(3, col).value for col in (7, 8)) == metric_headers:
        header_rows = 3
        expected_rate = "超时率"
        if ws.cell(1, 7).value != "网点交件":
            raise ValueError("T-1 新版模板缺少网点交件表头")
    else:
        raise ValueError("T-1 文件指标表头不符合已知模板")

    if ws.cell(header_rows, 9).value != expected_rate:
        raise ValueError("T-1 文件超时率列不符合已知模板")
    group_row = header_rows - 1
    expected_groups = {
        7: "揽收->入首分拨(工单列)",
        8: "揽收->入首分拨(超时列)",
        10: "揽收->入首分拨(第三列)",
        11: "揽收->入首分拨(第四列)",
        12: "揽收->入首分拨(第五列)",
        13: "揽收->入首分拨(第六列)",
    }
    if any(ws.cell(group_row, col).value != label for col, label in expected_groups.items()):
        raise ValueError("T-1 文件交件指标列顺序不符合已知模板")
    return header_rows


def download_and_save(fname, url):
    with urllib.request.urlopen(url, timeout=120) as response:
        wb = openpyxl.load_workbook(io.BytesIO(response.read()))
    ws = wb.active
    header_rows = header_row_count(ws)
    new_wb = openpyxl.Workbook()
    new_ws = new_wb.active
    max_row = min(header_rows + KEEP_DATA_ROWS, ws.max_row)
    for row_idx in range(1, max_row + 1):
        for cell in ws[row_idx]:
            new_cell = new_ws.cell(row=row_idx, column=cell.column, value=cell.value)
            copy_cell_style(cell, new_cell)
    for col_letter, dim in ws.column_dimensions.items():
        new_ws.column_dimensions[col_letter].width = dim.width
    for row_num in range(1, max_row + 1):
        if row_num in ws.row_dimensions:
            new_ws.row_dimensions[row_num].height = ws.row_dimensions[row_num].height
    for mc in ws.merged_cells.ranges:
        if mc.min_row <= max_row and mc.max_row <= max_row:
            new_ws.merge_cells(str(mc))
    new_wb.save(OUT_DIR / fname)
    wb.close()
    return new_ws.max_row


def main() -> int:
    print("=" * 50)
    print("  交件超时数据源更新工具")
    print("=" * 50)
    if not LINK_FILE.is_file():
        print(f"未找到链接清单: {LINK_FILE}")
        return 2
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    links = read_links(LINK_FILE)
    if not links:
        print("链接清单中没有有效的下载链接。")
        return 2

    print(f"共 {len(links)} 个文件，每个保留表头及前 {KEEP_DATA_ROWS} 行正文...\n")
    ok = fail = 0
    for i, (fname, url) in enumerate(links, 1):
        try:
            print(f"[{i}/{len(links)}] {fname} ...", end=" ", flush=True)
            rows = download_and_save(fname, url)
            print(f"OK ({rows}行)")
            ok += 1
        except Exception as e:
            print(f"FAIL: {e}")
            fail += 1

    print(f"\n下载完成: 成功 {ok}, 失败 {fail} / 共 {len(links)}")

    if fail:
        print("下载存在失败，未刷新看板数据。")
        return 1
    print("下载完成，等待 BAT 入口刷新看板数据。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
