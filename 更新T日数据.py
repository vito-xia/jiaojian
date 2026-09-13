"""根据 T 日内网轨迹生成当天的抖音 / 淘宝临时监控数据。

内网轨迹是 T 日的主数据源：按唯一运单计算待发运和始发退回时长，链接清单
下载的 Excel 只作为可选的客户归属补充。脚本不会把称重、装车或轨迹文本当作
实际交件时间，也不会修改常规 T-1/T-2 数据；运行完成后只写入当天的
``data/dashboard_tday.js``。
"""
from __future__ import annotations

import argparse
import hashlib
import math
import re
import shutil
import tempfile
import urllib.request
from collections import Counter, defaultdict
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from process_data import customer_is_excluded, parent_of, province_of, read_mapping, text, write_js_payload


TDAY_PLATFORMS = ("抖音", "淘宝")
ROLES = ("summary", "detail")
ROLE_LABELS = {"summary": "汇总", "detail": "TOP客户明细"}
EMPTY_VALUES = {"", "-", "--", "—", "/", "无", "暂无", "null", "none", "nan"}

SUMMARY_ALIASES: dict[str, tuple[str, ...]] = {
    "rank": ("排名", "序号", "rank"),
    "branch": ("分部名称", "网点名称", "分部", "网点", "站点"),
    "customer": ("客户名称", "商家名称", "客户", "商家"),
    "customer_code": ("客户编码", "客户代码", "商家编码", "客户id", "客户编号"),
    "timeout_24h": (
        "24H超时量", "24h超时量", "24H超时单量", "24h超时单量", "24H超时件数",
        "24h超时件数", "24小时超时量", "24小时超时单量", "24小时超时件数", "24H量", "24h量",
    ),
    "timeout_36h": (
        "36H超时量", "36h超时量", "36H超时单量", "36h超时单量", "36H超时件数",
        "36h超时件数", "36小时超时量", "36小时超时单量", "36小时超时件数", "36H量", "36h量",
    ),
}

DETAIL_ALIASES: dict[str, tuple[str, ...]] = {
    "branch": SUMMARY_ALIASES["branch"],
    "customer": SUMMARY_ALIASES["customer"],
    "customer_code": SUMMARY_ALIASES["customer_code"],
    "platform": ("平台", "平台名称", "来源平台"),
    "waybill": ("运单号", "快递单号", "物流单号", "订单号", "包裹号", "面单号", "waybill", "tracking"),
    "event_type": ("操作类型", "扫描类型", "轨迹类型", "节点类型", "事件类型", "操作名称", "节点名称", "节点", "轨迹", "扫描内容"),
    "event_time": ("操作时间", "扫描时间", "轨迹时间", "发生时间", "记录时间", "事件时间", "时间"),
    "pickup_time": ("首条揽收时间", "首次揽收时间", "第一条揽收时间", "揽收时间", "揽件时间", "首次揽件时间"),
    "handover_time": ("实际交件时间", "首条实际交件时间", "首次实际交件时间", "分拨交件时间", "首条分拨交件时间", "首次分拨交件时间", "交件时间"),
    "weighing_time": ("首条称重时间", "首次称重时间", "第一条称重时间", "称重时间", "称重扫描时间"),
    "snapshot_time": ("数据时间", "统计时间", "导出时间", "更新时间", "截止时间", "快照时间"),
}

# T 日内网轨迹是实际交件时长的唯一口径：第一条揽收时间到第一条分拨扫描时间。
# 这里单独维护别名，避免把普通明细里的“称重”或其他分拨节点误当成实际交件。
TRACE_ALIASES: dict[str, tuple[str, ...]] = {
    "waybill": DETAIL_ALIASES["waybill"],
    "pickup_time": ("第一条揽收时间", "首条揽收时间", "首次揽收时间", "揽收时间"),
    "handover_time": ("第一条分拨扫描时间", "首条分拨扫描时间", "首次分拨扫描时间"),
    "pickup_branch": ("第一条揽收网点名称", "首条揽收网点名称", "发件网点名称", "揽收网点名称"),
    "dispatch_branch": ("第一条分拨名称", "首条分拨名称", "分拨名称"),
    # “最新状态”与“最新轨迹状态”是两列不同语义；始发退回判定必须使用前者。
    "latest_status": ("最新状态", "当前状态", "状态"),
    "latest_scan_time": ("最新扫描时间", "最新轨迹时间", "最后扫描时间"),
    "exception_time": ("第一条异常记录时间", "首条异常记录时间", "首次异常记录时间"),
    "source": ("录单全部来源", "来源", "平台来源"),
}


class TDayFormatError(RuntimeError):
    """源表存在但无法识别必要字段。"""


def normalize(value: Any) -> str:
    return re.sub(r"[\s_\-—–·/\\（）()【】\[\]：:,.，。]+", "", text(value)).lower()


def normalize_waybill(value: Any) -> str:
    """将 Excel 中可能被读成数字/带空格的单号统一成可关联的键。"""
    raw = text(value).strip()
    if not raw:
        return ""
    if re.fullmatch(r"\d+\.0+", raw):
        raw = raw.split(".", 1)[0]
    return re.sub(r"\s+", "", raw).upper()


def safe_error_message(error: Exception) -> str:
    """错误日志中隐藏可能带权限参数的下载地址。"""
    return re.sub(r"https?://\S+", "[下载地址]", str(error), flags=re.IGNORECASE)


def optional_number(value: Any) -> int | float | None:
    raw = text(value)
    if raw.lower() in EMPTY_VALUES:
        return None
    cleaned = raw.replace(",", "").replace("，", "").replace("%", "")
    try:
        amount = float(cleaned)
    except (TypeError, ValueError):
        return None
    return int(round(amount)) if amount.is_integer() else round(amount, 4)


def row_value(row: Iterable[Any], index: int | None) -> Any:
    if index is None:
        return None
    values = tuple(row)
    return values[index] if index < len(values) else None


def alias_score(value: Any, alias: str) -> int:
    candidate = normalize(value)
    target = normalize(alias)
    if not candidate or not target:
        return 0
    if candidate == target:
        return 100 + len(target)
    if target in candidate:
        return 60 + len(target)
    return 0


def detect_columns(
    rows: list[tuple[Any, ...]],
    aliases: dict[str, tuple[str, ...]],
    required: tuple[str, ...],
    header_limit: int = 12,
) -> tuple[int, dict[str, int]] | None:
    """识别单行或多级表头，返回数据起始前一行和字段列号。"""
    if not rows:
        return None
    max_columns = max(len(row) for row in rows)
    best: tuple[int, int, int, dict[str, int]] | None = None
    for end in range(min(header_limit, len(rows))):
        combined = []
        for column in range(max_columns):
            combined.append("".join(text(rows[row][column]) for row in range(end + 1) if column < len(rows[row])))
        columns: dict[str, int] = {}
        used: set[int] = set()
        score_total = 0
        ordered_fields = list(required) + [field for field in aliases if field not in required]
        for field in ordered_fields:
            candidates = sorted(
                ((alias_score(combined[column], alias), column) for column in range(max_columns) for alias in aliases.get(field, ())),
                reverse=True,
            )
            for score, column in candidates:
                if score and column not in used:
                    columns[field] = column
                    used.add(column)
                    score_total += score
                    break
        required_count = sum(field in columns for field in required)
        if required_count < len(required):
            continue
        candidate = (required_count, score_total, -end, columns)
        if best is None or candidate[:3] > best[:3]:
            best = candidate
    if best is None:
        return None
    return -best[2], best[3]


def parse_datetime(value: Any, epoch: Any, default_day: date) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch)
            return converted.replace(tzinfo=None) if isinstance(converted, datetime) else datetime.combine(converted, datetime.min.time())
        except (TypeError, ValueError, OverflowError):
            return None
    raw = text(value)
    if raw.lower() in EMPTY_VALUES:
        return None
    raw = raw.replace("T", " ").replace("年", "-").replace("月", "-").replace("日", " ").replace("/", "-").replace(".", "-")
    raw = re.sub(r"\s+", " ", raw).strip()
    for fmt in (
        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%m-%d %H:%M:%S", "%m-%d %H:%M",
        "%Y%m%d%H%M%S", "%Y%m%d%H%M", "%Y%m%d",
    ):
        try:
            parsed = datetime.strptime(raw[:19], fmt)
            if parsed.year == 1900:
                parsed = parsed.replace(year=default_day.year)
            return parsed
        except ValueError:
            continue
    match = re.fullmatch(r"(\d{1,2})\s*[-:]\s*(\d{1,2})(?::(\d{1,2}))?", raw)
    if match:
        return datetime(default_day.year, default_day.month, default_day.day, int(match.group(1)), int(match.group(2)), int(match.group(3) or 0))
    return None


def parse_explicit_day(value: str, default_year: int) -> date | None:
    iso = re.search(r"(?<!\d)(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})(?!\d)", value)
    if iso:
        try:
            return date(int(iso.group(1)), int(iso.group(2)), int(iso.group(3)))
        except ValueError:
            return None
    cn = re.search(r"(?<!\d)(\d{1,2})月(\d{1,2})日", value)
    if cn:
        try:
            return date(default_year, int(cn.group(1)), int(cn.group(2)))
        except ValueError:
            return None
    return None


def parse_link_role(label: str) -> str | None:
    normalized = normalize(label)
    if "汇总" in normalized or "总表" in normalized or "summary" in normalized:
        return "summary"
    if "top客户" in normalized or "客户明细" in normalized or "明细" in normalized or "detail" in normalized:
        return "detail"
    return None


def parse_links(path: Path, target_day: date) -> tuple[dict[tuple[str, str], dict[str, Any]], list[str]]:
    entries: dict[tuple[str, str], dict[str, Any]] = {}
    untyped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    warnings: list[str] = []
    if not path.exists():
        return entries, [f"未找到链接清单：{path}"]
    for line_number, raw_line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1):
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        parts = line.split()
        if len(parts) < 2:
            warnings.append(f"第 {line_number} 行无法解析")
            continue
        url = parts[-1]
        label = " ".join(parts[:-1])
        platform = next((item for item in TDAY_PLATFORMS if item in label or (item == "淘宝" and "淘天" in label)), None)
        source_day = parse_explicit_day(label, target_day.year)
        if not platform:
            continue
        role = parse_link_role(label)
        if not role:
            if source_day in (None, target_day):
                untyped[platform].append({"url": url, "label": label, "source_day": source_day, "line_number": line_number})
            continue
        if source_day and source_day != target_day:
            warnings.append(f"第 {line_number} 行不是目标 T 日，已跳过")
            continue
        key = (platform, role)
        if key in entries:
            warnings.append(f"第 {line_number} 行重复定义 {platform}{ROLE_LABELS[role]}，已使用最后一条")
        entries[key] = {"url": url, "label": label, "source_day": source_day}
    # 兼容仅写“平台 + 日期 + 链接”的简写：同一平台当天的第一条按汇总、
    # 第二条按 TOP 客户明细处理。若用户已经在标签中写明角色，则优先使用显式角色。
    for platform, candidates in untyped.items():
        for index, entry in enumerate(candidates[: len(ROLES)]):
            role = ROLES[index]
            key = (platform, role)
            if key in entries:
                warnings.append(f"第 {entry['line_number']} 行未标注角色，因已有显式 {ROLE_LABELS[role]} 链接而跳过")
                continue
            entries[key] = {key_name: entry[key_name] for key_name in ("url", "label", "source_day")}
            warnings.append(f"第 {entry['line_number']} 行未标注角色，已按第 {index + 1} 条顺序识别为{ROLE_LABELS[role]}")
        if len(candidates) > len(ROLES):
            warnings.append(f"{platform} 当天无角色链接超过 {len(ROLES)} 条，后续链接已跳过")
    return entries, warnings


def read_summary(path: Path, platform: str, target_day: date, mapping: dict[str, dict[str, str]]) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    selected_rows: list[tuple[Any, ...]] | None = None
    header_end = 0
    columns: dict[str, int] = {}
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        detected = detect_columns(rows[:12], SUMMARY_ALIASES, ("branch", "customer", "timeout_24h", "timeout_36h"))
        if not detected:
            # 常规交件汇总表的多级表头把 24H/36H 写成“票件量A / 超时量B”，
            # 与历史 T-1 源表保持同一列顺序：分部、客户后第 7/8 列即为两项量。
            legacy = detect_columns(rows[:12], SUMMARY_ALIASES, ("branch", "customer"))
            if legacy and max(len(row) for row in rows) >= 8:
                legacy_end, legacy_columns = legacy
                legacy_columns["timeout_24h"] = 6
                legacy_columns["timeout_36h"] = 7
                detected = legacy_end, legacy_columns
        if detected:
            header_end, columns = detected
            selected_rows = rows
            break
    workbook.close()
    if selected_rows is None:
        raise TDayFormatError("未识别汇总表的网点、客户、24H或36H字段")

    rows: list[dict[str, Any]] = []
    for source_index, row in enumerate(selected_rows[header_end + 1 :], 1):
        branch = text(row_value(row, columns.get("branch")))
        customer = text(row_value(row, columns.get("customer")))
        if not branch or not customer or branch in {"合计", "总计", "汇总"} or customer in {"合计", "总计", "汇总"}:
            continue
        if customer_is_excluded(customer):
            continue
        customer_code = text(row_value(row, columns.get("customer_code")))
        identity = "|".join((platform, branch, customer_code or customer, str(source_index)))
        row_id = f"{platform}-{hashlib.sha1(identity.encode('utf-8')).hexdigest()[:14]}"
        rows.append({
            "id": row_id,
            "rank": len(rows) + 1,
            "platform": platform,
            "date": target_day.isoformat(),
            "branch": branch,
            "parent_name": parent_of(branch, mapping),
            "province": province_of(branch, mapping),
            "customer": customer,
            "customer_code": customer_code,
            "timeout_24h": optional_number(row_value(row, columns.get("timeout_24h"))),
            "timeout_36h": optional_number(row_value(row, columns.get("timeout_36h"))),
        })
        if len(rows) >= 10:
            break
    return {
        "status": "uploaded",
        "source_rows": len(selected_rows) - header_end - 1,
        "rows": rows,
        "error": "",
    }


@lru_cache(maxsize=2)
def read_trace_source(path: Path, target_day: date, fallback_snapshot: datetime) -> dict[str, Any]:
    """读取 T 日内网轨迹，并按唯一运单保留当前状态和首个事件。

    已出现第一条分拨扫描的运单不进入待发运曲线；没有分拨扫描且最新状态
    不是“始发退回”的运单按数据截止时间计算待发运时长。始发退回单独按
    ``第一条异常记录时间 - 第一条揽收时间``计算。称重字段和轨迹文本不参与。
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: dict[str, dict[str, Any]] = {}
    source_rows = 0
    invalid_rows = 0
    pickup_rows = 0
    handover_rows = 0
    pending_rows = 0
    return_rows = 0
    return_valid_rows = 0
    header_end = 0
    columns: dict[str, int] = {}
    selected_sheet = None
    max_event: datetime | None = None
    for sheet in workbook.worksheets:
        scan_rows: list[tuple[Any, ...]] = []
        iterator = sheet.iter_rows(values_only=True)
        for _ in range(12):
            try:
                scan_rows.append(next(iterator))
            except StopIteration:
                break
        detected = detect_columns(scan_rows, TRACE_ALIASES, ("waybill", "pickup_time", "handover_time"))
        if detected:
            header_end, columns = detected
            selected_sheet = sheet
            break
    if selected_sheet is None:
        workbook.close()
        raise TDayFormatError("未识别内网轨迹的运单号、第一条揽收时间或第一条分拨扫描时间字段")

    iterator = selected_sheet.iter_rows(values_only=True)
    # 重新创建迭代器，跳过表头行；避免把前 12 行缓存重复计数。
    for _ in range(header_end + 1):
        try:
            next(iterator)
        except StopIteration:
            break
    for row in iterator:
        source_rows += 1
        waybill = normalize_waybill(row_value(row, columns.get("waybill")))
        if not waybill:
            invalid_rows += 1
            continue
        pickup = parse_datetime(row_value(row, columns.get("pickup_time")), workbook.epoch, target_day)
        handover = parse_datetime(row_value(row, columns.get("handover_time")), workbook.epoch, target_day)
        exception_time = parse_datetime(row_value(row, columns.get("exception_time")), workbook.epoch, target_day)
        latest_scan = parse_datetime(row_value(row, columns.get("latest_scan_time")), workbook.epoch, target_day)
        latest_status = text(row_value(row, columns.get("latest_status")))
        source = text(row_value(row, columns.get("source")))
        for event_time in (pickup, handover, exception_time, latest_scan):
            if event_time is not None and (max_event is None or event_time > max_event):
                max_event = event_time
        if pickup is None:
            invalid_rows += 1
            continue
        pickup_rows += 1
        record = records.setdefault(waybill, {
            "pickup": None,
            "handover": None,
            "exception_time": None,
            "latest_scan": None,
            "latest_status": "",
            "source_values": set(),
            "pickup_branch": text(row_value(row, columns.get("pickup_branch"))),
            "dispatch_branch": text(row_value(row, columns.get("dispatch_branch"))),
        })
        update_min(record, "pickup", pickup)
        update_min(record, "handover", handover)
        update_min(record, "exception_time", exception_time)
        if source:
            record["source_values"].add(source)
        if not record.get("pickup_branch"):
            record["pickup_branch"] = text(row_value(row, columns.get("pickup_branch")))
        if not record.get("dispatch_branch"):
            record["dispatch_branch"] = text(row_value(row, columns.get("dispatch_branch")))
        if latest_scan is not None and (record.get("latest_scan") is None or latest_scan >= record["latest_scan"]):
            record["latest_scan"] = latest_scan
            record["latest_status"] = latest_status
        elif record.get("latest_scan") is None and not record.get("latest_status"):
            record["latest_status"] = latest_status
    workbook.close()

    snapshot = max_event or fallback_snapshot
    for record in records.values():
        pickup = record.get("pickup")
        handover = record.get("handover")
        status = record.get("latest_status") or ""
        if handover is not None:
            elapsed = (handover - pickup).total_seconds() / 3600
            if elapsed < 0:
                record["invalid"] = True
                invalid_rows += 1
                continue
            record["elapsed"] = round(elapsed, 3)
            record["pending"] = False
            handover_rows += 1
        elif status == "始发退回":
            record["pending"] = False
            return_rows += 1
        else:
            elapsed = (snapshot - pickup).total_seconds() / 3600
            if elapsed >= 0:
                record["elapsed"] = round(elapsed, 3)
                record["pending"] = True
                pending_rows += 1
            else:
                record["invalid"] = True
                invalid_rows += 1
        if status == "始发退回":
            exception_time = record.get("exception_time")
            if exception_time is not None:
                return_elapsed = (exception_time - pickup).total_seconds() / 3600
                if return_elapsed >= 0:
                    record["return_elapsed"] = round(return_elapsed, 3)
                    return_valid_rows += 1
                else:
                    record["invalid_return"] = True
    return {
        "status": "uploaded",
        "records": records,
        "source_rows": source_rows,
        "waybill_count": len(records),
        "duplicate_rows": max(0, source_rows - len(records)),
        "pickup_rows": pickup_rows,
        "handover_rows": handover_rows,
        "pending_rows": pending_rows,
        "return_rows": return_rows,
        "return_valid_rows": return_valid_rows,
        "invalid_rows": invalid_rows,
        "snapshot_at": snapshot,
        "source_day": max_event.date().isoformat() if max_event else "",
        "freshness_status": "uploaded" if datetime.fromtimestamp(path.stat().st_mtime).date() == target_day or (max_event and max_event.date() == target_day) else "stale",
    }


def target_indexes(rows: list[dict[str, Any]]) -> dict[str, dict[Any, list[str]]]:
    indexes = {"code": defaultdict(list), "branch_customer": defaultdict(list), "customer": defaultdict(list)}
    for row in rows:
        row_id = row["id"]
        code = normalize(row.get("customer_code"))
        branch_customer = (normalize(row.get("branch")), normalize(row.get("customer")))
        customer = normalize(row.get("customer"))
        if code:
            indexes["code"][code].append(row_id)
        if branch_customer[0] and branch_customer[1]:
            indexes["branch_customer"][branch_customer].append(row_id)
        if customer:
            indexes["customer"][customer].append(row_id)
    return indexes


def match_detail_targets(branch: str, customer: str, code: str, indexes: dict[str, dict[Any, list[str]]]) -> list[str]:
    branch_key = normalize(branch)
    customer_key = normalize(customer)
    code_key = normalize(code)
    by_branch = indexes["branch_customer"].get((branch_key, customer_key), [])
    if len(by_branch) == 1:
        return by_branch
    by_code = indexes["code"].get(code_key, []) if code_key else []
    if len(by_code) == 1:
        return by_code
    by_customer = indexes["customer"].get(customer_key, [])
    return by_customer if len(by_customer) == 1 else []


def update_min(group: dict[str, datetime | None], key: str, value: datetime | None) -> None:
    if value is None:
        return
    if group.get(key) is None or value < group[key]:
        group[key] = value


def read_detail(
    path: Path,
    platform: str,
    summary_rows: list[dict[str, Any]],
    target_day: date,
    fallback_snapshot: datetime,
    trace_records: dict[str, dict[str, Any]] | None = None,
    trace_only: bool = False,
) -> dict[str, Any]:
    indexes = target_indexes(summary_rows)
    groups: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    snapshot_by_target: dict[str, datetime] = {}
    matched_rows = 0
    invalid_rows = 0
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        iterator = sheet.iter_rows(values_only=True)
        scan_rows: list[tuple[Any, ...]] = []
        for _ in range(12):
            try:
                scan_rows.append(next(iterator))
            except StopIteration:
                break
        detected = detect_columns(scan_rows, DETAIL_ALIASES, ("event_time",))
        if not detected:
            detected = detect_columns(scan_rows, DETAIL_ALIASES, ("pickup_time", "handover_time"))
        if not detected:
            # TOP 客户明细也可能是一票一行的关联表，只提供网点/客户/单号，
            # 实际时间另由 T 日内网轨迹补充；此时仍需先建立单号归属关系。
            detected = detect_columns(scan_rows, DETAIL_ALIASES, ("waybill",))
        if not detected:
            continue
        header_end, columns = detected
        def process_row(row_number: int, row: tuple[Any, ...]) -> None:
            nonlocal matched_rows, invalid_rows
            branch = text(row_value(row, columns.get("branch")))
            customer = text(row_value(row, columns.get("customer")))
            code = text(row_value(row, columns.get("customer_code")))
            target_ids = match_detail_targets(branch, customer, code, indexes)
            if len(target_ids) != 1:
                return
            target_id = target_ids[0]
            matched_rows += 1
            waybill = normalize_waybill(row_value(row, columns.get("waybill"))) or f"{sheet.title}:{row_number}"
            group = groups[target_id].setdefault(waybill, {"pickup": None, "handover": None, "weighing": None})
            snapshot = parse_datetime(row_value(row, columns.get("snapshot_time")), workbook.epoch, target_day)
            if snapshot and (target_id not in snapshot_by_target or snapshot > snapshot_by_target[target_id]):
                snapshot_by_target[target_id] = snapshot

            pickup = parse_datetime(row_value(row, columns.get("pickup_time")), workbook.epoch, target_day)
            handover = parse_datetime(row_value(row, columns.get("handover_time")), workbook.epoch, target_day)
            weighing = parse_datetime(row_value(row, columns.get("weighing_time")), workbook.epoch, target_day)
            if pickup is not None or handover is not None or weighing is not None:
                update_min(group, "pickup", pickup)
                update_min(group, "handover", handover)
                update_min(group, "weighing", weighing)
                return

            # 关联表没有事件时间时，单号已记录，等待下方与内网轨迹合并。
            if columns.get("event_type") is None and columns.get("event_time") is None:
                return

            event = text(row_value(row, columns.get("event_type")))
            event_time = parse_datetime(row_value(row, columns.get("event_time")), workbook.epoch, target_day)
            if not event or event_time is None:
                invalid_rows += 1
                return
            if "揽收" in event or "揽件" in event:
                update_min(group, "pickup", event_time)
            if "实际交件" in event or "分拨交件" in event or "交件" in event:
                update_min(group, "handover", event_time)
            if "称重" in event:
                update_min(group, "weighing", event_time)
        rows_before_header = scan_rows[header_end + 1 :]
        for row_number, row in enumerate(rows_before_header, header_end + 2):
            process_row(row_number, row)
        for row_number, row in enumerate(iterator, header_end + 2 + len(rows_before_header)):
            process_row(row_number, row)
    workbook.close()

    if trace_only:
        # TOP 明细在 T 日链路中只负责“客户/网点 ↔ 单号”归属。任何随表携带的
        # 称重、事件或其他时间都先清空，最终时长必须来自内网轨迹固定三字段。
        for by_waybill in groups.values():
            for group in by_waybill.values():
                group.update({"pickup": None, "handover": None, "weighing": None})

    trace_matched = 0
    if trace_records:
        for by_waybill in groups.values():
            for waybill, group in by_waybill.items():
                trace = trace_records.get(normalize_waybill(waybill))
                if not trace:
                    continue
                trace_matched += 1
                update_min(group, "pickup", trace.get("pickup"))
                update_min(group, "handover", trace.get("handover"))
                if trace.get("pending") and trace.get("elapsed") is not None:
                    group["pending"] = True
                    group["trace_elapsed"] = trace.get("elapsed")

    result: dict[str, Any] = {}
    completed_count = 0
    pending_total = 0
    for row in summary_rows:
        target_id = row["id"]
        durations: list[float] = []
        pending_count = 0
        invalid_count = 0
        for group in groups.get(target_id, {}).values():
            pickup = group.get("pickup")
            handover = group.get("handover")
            if pickup is None:
                invalid_count += 1
                continue
            if handover is not None:
                elapsed = (handover - pickup).total_seconds() / 3600
                if elapsed < 0:
                    invalid_count += 1
                    continue
                durations.append(round(elapsed, 3))
                completed_count += 1
                continue
            pending_elapsed = group.get("trace_elapsed")
            if group.get("pending") and isinstance(pending_elapsed, (int, float)) and pending_elapsed >= 0:
                durations.append(round(float(pending_elapsed), 3))
                pending_count += 1
                pending_total += 1
                continue
            invalid_count += 1
        result[target_id] = {
            "durations": durations,
            "pending_count": pending_count,
            "invalid_count": invalid_count,
            "shipment_count": len(groups.get(target_id, {})),
            "matched_rows": matched_rows,
            "invalid_rows": invalid_rows,
        }
    has_handover = completed_count > 0
    has_pending = pending_total > 0
    return {
        "status": "uploaded",
        "handover_available": has_handover,
        "pending_available": has_pending,
        "rows": result,
        "matched_rows": matched_rows,
        "invalid_rows": invalid_rows,
        "trace_matched": trace_matched,
    }


def survival_points(durations: list[float], threshold: int) -> list[dict[str, int]]:
    clean = [value for value in durations if math.isfinite(value) and value >= 0]
    if not clean:
        return []
    end = max(threshold + 1, math.ceil(max(clean)) + 1)
    span = end - threshold
    step = max(1, math.ceil(span / 360))
    hours = list(range(threshold, end + 1, step))
    if hours[-1] != end:
        hours.append(end)
    return [{"hours": hour, "count": sum(value >= hour for value in clean)} for hour in hours]


def survival_points_from_zero(durations: list[float]) -> list[dict[str, int]]:
    """生成从 0H 开始的非递增累计曲线点。"""
    clean = sorted(value for value in durations if math.isfinite(value) and value >= 0)
    if not clean:
        return []
    end = max(36, math.ceil(clean[-1]))
    step = max(1, math.ceil(end / 360))
    hours = list(range(0, end + 1, step))
    if hours[-1] != end:
        hours.append(end)
    return [{"hours": hour, "count": sum(value >= hour for value in clean)} for hour in hours]


def classify_platforms(values: Iterable[Any]) -> set[str]:
    """从来源文本中识别抖音/淘宝；同时出现时返回两个平台。"""
    raw = " ".join(text(value) for value in values if text(value))
    normalized = normalize(raw)
    platforms: set[str] = set()
    if any(token in normalized for token in ("抖音", "douyin", "tiktok")):
        platforms.add("抖音")
    if any(token in normalized for token in ("淘宝", "淘天", "taobao", "cainiao", "菜鸟")):
        platforms.add("淘宝")
    return platforms


def most_common_value(counter: Counter[str] | None) -> str:
    if not counter:
        return ""
    return counter.most_common(1)[0][0]


def read_optional_waybill_metadata(source_dir: Path) -> tuple[dict[str, dict[str, Any]], dict[str, Any]]:
    """扫描下载目录中的任意 Excel，尽力建立单号到客户/网点的补充映射。

    这些文件不参与时间计算；无法识别的文件只进入质量信息，不阻塞内网轨迹主链路。
    """
    metadata: dict[str, dict[str, Any]] = {}
    files = sorted(path for path in source_dir.glob("*.xlsx") if path.is_file())
    recognized_files = 0
    metadata_rows = 0
    errors: list[str] = []
    for path in files:
        recognized_in_file = False
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            for sheet in workbook.worksheets:
                iterator = sheet.iter_rows(values_only=True)
                scan_rows: list[tuple[Any, ...]] = []
                for _ in range(12):
                    try:
                        scan_rows.append(next(iterator))
                    except StopIteration:
                        break
                detected = detect_columns(scan_rows, DETAIL_ALIASES, ("waybill",))
                if not detected:
                    continue
                recognized_in_file = True
                header_end, columns = detected
                rows_before_header = scan_rows[header_end + 1 :]
                def process_metadata_row(row: tuple[Any, ...]) -> None:
                    nonlocal metadata_rows
                    waybill = normalize_waybill(row_value(row, columns.get("waybill")))
                    if not waybill:
                        return
                    item = metadata.setdefault(waybill, {
                        "customers": Counter(),
                        "customer_codes": Counter(),
                        "branches": Counter(),
                        "platforms": Counter(),
                    })
                    customer = text(row_value(row, columns.get("customer")))
                    code = text(row_value(row, columns.get("customer_code")))
                    branch = text(row_value(row, columns.get("branch")))
                    platform_hint = text(row_value(row, columns.get("platform")))
                    if customer and not customer_is_excluded(customer):
                        item["customers"][customer] += 1
                    if code:
                        item["customer_codes"][code] += 1
                    if branch:
                        item["branches"][branch] += 1
                    for platform in classify_platforms((platform_hint, path.stem, sheet.title)):
                        item["platforms"][platform] += 1
                    metadata_rows += 1
                for _, row in enumerate(rows_before_header, header_end + 2):
                    process_metadata_row(row)
                for _, row in enumerate(iterator, header_end + 2 + len(rows_before_header)):
                    process_metadata_row(row)
            workbook.close()
        except Exception as error:
            errors.append(f"{path.name}：{safe_error_message(error)}")
        if recognized_in_file:
            recognized_files += 1
    return metadata, {
        "file_count": len(files),
        "recognized_file_count": recognized_files,
        "waybill_count": len(metadata),
        "metadata_rows": metadata_rows,
        "errors": errors,
    }


def empty_trace_platform(platform: str, status: str, error: str = "") -> tuple[dict[str, Any], dict[str, Any]]:
    return {
        "source_status": status,
        "summary_status": status,
        "detail_status": "not_required",
        "summary_rows": 0,
        "rows": [],
        "error": error,
        "row_dimension": "branch",
        "customer_mapping_status": "not_available",
        "kpis": {},
    }, {}


def build_trace_platform_payload(
    platform: str,
    trace: dict[str, Any],
    metadata: dict[str, dict[str, Any]],
    metadata_report: dict[str, Any],
    mapping: dict[str, dict[str, str]],
    target_day: date,
) -> tuple[dict[str, Any], dict[str, dict[str, Any]]]:
    if trace.get("freshness_status") != "uploaded":
        return empty_trace_platform(platform, "not_uploaded", "T日内网轨迹不是目标日期的当前快照")

    groups: dict[tuple[str, str, str], dict[str, Any]] = {}
    platform_records = 0
    mapped_waybills = 0
    mixed_waybills = 0
    unclassified_waybills = 0
    pending_durations: list[float] = []
    return_durations: list[float] = []

    for waybill, record in trace.get("records", {}).items():
        if record.get("invalid"):
            continue
        candidates = classify_platforms(record.get("source_values", ()))
        meta = metadata.get(waybill, {})
        hinted = set((meta.get("platforms") or Counter()).keys())
        if len(candidates) != 1 and len(hinted) == 1:
            candidates = hinted
        if len(candidates) > 1:
            mixed_waybills += 1
            continue
        if not candidates:
            unclassified_waybills += 1
            continue
        resolved_platform = next(iter(candidates))
        if resolved_platform != platform:
            continue
        platform_records += 1

        customer = most_common_value(meta.get("customers"))
        customer_code = most_common_value(meta.get("customer_codes"))
        branch = text(record.get("pickup_branch")) or text(record.get("dispatch_branch")) or "未标注网点"
        if customer:
            mapped_waybills += 1
        group_key = (branch, customer, customer_code)
        group = groups.setdefault(group_key, {
            "branch": branch,
            "customer": customer,
            "customer_code": customer_code,
            "pending": [],
            "start_return": [],
            "waybill_count": 0,
        })
        group["waybill_count"] += 1
        if record.get("pending") and isinstance(record.get("elapsed"), (int, float)):
            value = float(record["elapsed"])
            group["pending"].append(value)
            pending_durations.append(value)
        if record.get("latest_status") == "始发退回" and isinstance(record.get("return_elapsed"), (int, float)):
            value = float(record["return_elapsed"])
            group["start_return"].append(value)
            return_durations.append(value)

    def sort_key(group: dict[str, Any]) -> tuple[Any, ...]:
        return (
            -sum(value >= 36 for value in group["pending"]),
            -sum(value >= 24 for value in group["pending"]),
            -len(group["pending"]),
            text(group["branch"]),
            text(group["customer"]),
        )

    sorted_groups = sorted(groups.values(), key=sort_key)
    rows: list[dict[str, Any]] = []
    trend_map: dict[str, dict[str, Any]] = {}
    for rank, group in enumerate(sorted_groups[:10], 1):
        identity = "|".join((platform, group["branch"], group["customer_code"] or group["customer"], str(rank)))
        row_id = f"{platform}-{hashlib.sha1(identity.encode('utf-8')).hexdigest()[:14]}"
        pending = group["pending"]
        start_return = group["start_return"]
        row = {
            "id": row_id,
            "rank": rank,
            "platform": platform,
            "date": target_day.isoformat(),
            "branch": group["branch"],
            "parent_name": parent_of(group["branch"], mapping),
            "province": province_of(group["branch"], mapping),
            "customer": group["customer"],
            "customer_code": group["customer_code"],
            "timeout_24h": sum(value >= 24 for value in pending),
            "timeout_36h": sum(value >= 36 for value in pending),
            "pending_count": len(pending),
            "start_return_count": len(start_return),
            "waybill_count": group["waybill_count"],
        }
        rows.append(row)
        trend_map[row_id] = {
            "summary_24h": row["timeout_24h"],
            "summary_36h": row["timeout_36h"],
            "detail_count": len(pending) + len(start_return),
            "shipment_count": group["waybill_count"],
            "pending_count": len(pending),
            "start_return_count": len(start_return),
            "invalid_count": 0,
            "trend_status": "ready" if pending or start_return else "no_data",
            "trends": {
                "pending": {"points": survival_points_from_zero(pending)},
                "start_return": {"points": survival_points_from_zero(start_return)},
            },
        }

    mapping_status = "ready" if platform_records and mapped_waybills == platform_records else "partial" if mapped_waybills else "not_available"
    platform_data = {
        "source_status": "uploaded",
        "summary_status": "uploaded",
        "detail_status": "uploaded" if metadata_report.get("recognized_file_count") else "not_required",
        "summary_rows": len(groups),
        "rows": rows,
        "error": "",
        "row_dimension": "customer" if mapped_waybills else "branch",
        "customer_mapping_status": mapping_status,
        "trace_source_status": "uploaded",
        "trace_source_rows": trace.get("source_rows", 0),
        "trace_waybill_count": trace.get("waybill_count", 0),
        "trace_duplicate_rows": trace.get("duplicate_rows", 0),
        "trace_handover_rows": trace.get("handover_rows", 0),
        "trace_pending_rows": trace.get("pending_rows", 0),
        "trace_return_rows": trace.get("return_rows", 0),
        "trace_return_valid_rows": trace.get("return_valid_rows", 0),
        "trace_snapshot_at": trace.get("snapshot_at").strftime("%Y-%m-%d %H:%M:%S") if trace.get("snapshot_at") else "",
        "trace_source_day": trace.get("source_day", ""),
        "mixed_waybills": mixed_waybills,
        "unclassified_waybills": unclassified_waybills,
        "kpis": {
            "total_waybills": platform_records,
            "pending_count": len(pending_durations),
            "pending_24h": sum(value >= 24 for value in pending_durations),
            "pending_36h": sum(value >= 36 for value in pending_durations),
            "start_return_count": len(return_durations),
            "start_return_24h": sum(value >= 24 for value in return_durations),
            "start_return_36h": sum(value >= 36 for value in return_durations),
            "group_count": len(groups),
            "mapped_waybills": mapped_waybills,
            "unmatched_waybills": max(0, platform_records - mapped_waybills),
        },
        "quality": {
            "metadata_file_count": metadata_report.get("file_count", 0),
            "metadata_recognized_file_count": metadata_report.get("recognized_file_count", 0),
            "metadata_waybill_count": metadata_report.get("waybill_count", 0),
            "metadata_errors": metadata_report.get("errors", []),
        },
    }
    return platform_data, trend_map


def build_trace_payload(
    source_dir: Path,
    target_day: date,
    mapping: dict[str, dict[str, str]],
    trace_path: Path | None,
    fallback_snapshot: datetime,
) -> tuple[dict[str, Any], dict[str, dict[str, Any]]]:
    if trace_path is None or not trace_path.exists():
        platforms = {platform: empty_trace_platform(platform, "not_uploaded", "T日内网轨迹未上传")[0] for platform in TDAY_PLATFORMS}
        return {
            "schema_version": 2,
            "date": target_day.isoformat(),
            "generated_at": fallback_snapshot.strftime("%Y-%m-%d %H:%M:%S"),
            "source_status": "not_uploaded",
            "platforms": platforms,
            "trends": {platform: {} for platform in TDAY_PLATFORMS},
        }, {platform: {} for platform in TDAY_PLATFORMS}
    try:
        trace = read_trace_source(trace_path, target_day, fallback_snapshot)
    except Exception as error:
        platforms = {platform: empty_trace_platform(platform, "invalid_format", str(error))[0] for platform in TDAY_PLATFORMS}
        return {
            "schema_version": 2,
            "date": target_day.isoformat(),
            "generated_at": fallback_snapshot.strftime("%Y-%m-%d %H:%M:%S"),
            "source_status": "invalid_format",
            "platforms": platforms,
            "trends": {platform: {} for platform in TDAY_PLATFORMS},
        }, {platform: {} for platform in TDAY_PLATFORMS}

    metadata, metadata_report = read_optional_waybill_metadata(source_dir)
    platforms: dict[str, Any] = {}
    trends: dict[str, dict[str, Any]] = {}
    for platform in TDAY_PLATFORMS:
        platforms[platform], trends[platform] = build_trace_platform_payload(
            platform, trace, metadata, metadata_report, mapping, target_day,
        )
    overall = "uploaded" if trace.get("freshness_status") == "uploaded" else "not_uploaded"
    return {
        "schema_version": 2,
        "date": target_day.isoformat(),
        "generated_at": fallback_snapshot.strftime("%Y-%m-%d %H:%M:%S"),
        "source_status": overall,
        "trace_source_status": trace.get("freshness_status", "not_uploaded"),
        "trace_source_rows": trace.get("source_rows", 0),
        "trace_waybill_count": trace.get("waybill_count", 0),
        "trace_duplicate_rows": trace.get("duplicate_rows", 0),
        "trace_snapshot_at": trace.get("snapshot_at").strftime("%Y-%m-%d %H:%M:%S") if trace.get("snapshot_at") else "",
        "trace_source_day": trace.get("source_day", ""),
        "platforms": platforms,
        "trends": trends,
    }, trends


def platform_payload(
    platform: str,
    summary_path: Path | None,
    detail_path: Path | None,
    trace_path: Path | None,
    target_day: date,
    mapping: dict[str, dict[str, str]],
    fallback_snapshot: datetime,
) -> tuple[dict[str, Any], dict[str, dict[str, Any]]]:
    if summary_path is None or not summary_path.exists():
        return {
            "source_status": "not_uploaded", "summary_status": "not_uploaded", "detail_status": "not_uploaded",
            "summary_rows": 0, "rows": [], "error": "汇总表未上传",
        }, {}
    try:
        summary = read_summary(summary_path, platform, target_day, mapping)
    except Exception as error:
        return {
            "source_status": "invalid_format", "summary_status": "invalid_format", "detail_status": "not_uploaded",
            "summary_rows": 0, "rows": [], "error": str(error),
        }, {}

    details: dict[str, Any] = {}
    trace_records: dict[str, dict[str, Any]] = {}
    trace_meta: dict[str, Any] = {
        "trace_source_status": "not_uploaded",
        "trace_source_rows": 0,
        "trace_waybill_count": 0,
        "trace_handover_rows": 0,
        "trace_pending_rows": 0,
        "trace_matched_waybills": 0,
    }
    if trace_path is not None and trace_path.exists():
        try:
            trace = read_trace_source(trace_path, target_day, fallback_snapshot)
            trace_records = trace["records"]
            trace_meta.update({
                "trace_source_status": "uploaded",
                "trace_source_rows": trace.get("source_rows", 0),
                "trace_waybill_count": trace.get("waybill_count", 0),
                "trace_handover_rows": trace.get("handover_rows", 0),
                "trace_pending_rows": trace.get("pending_rows", 0),
            })
        except Exception as error:
            trace_meta.update({"trace_source_status": "invalid_format", "trace_error": str(error)})
    detail_status = "not_uploaded"
    detail_error = "TOP客户明细未上传"
    if detail_path is not None and detail_path.exists():
        try:
            if trace_meta["trace_source_status"] != "uploaded":
                detail_status = "handover_missing"
                detail_error = (
                    f"内网轨迹格式异常：{trace_meta.get('trace_error', '')}"
                    if trace_meta["trace_source_status"] == "invalid_format"
                    else "T日内网轨迹未上传"
                )
            else:
                detail = read_detail(
                    detail_path, platform, summary["rows"], target_day, fallback_snapshot,
                    trace_records=trace_records,
                    trace_only=True,
                )
                details = detail["rows"]
                trace_meta["trace_matched_waybills"] = detail.get("trace_matched", 0)
                if detail.get("handover_available") or detail.get("pending_available"):
                    detail_status = "uploaded"
                    detail_error = ""
                else:
                    detail_status = "handover_missing"
                    detail_error = "内网轨迹已上传，但 TOP 客户明细未匹配到单号"
        except Exception as error:
            detail_status = "invalid_format"
            detail_error = str(error)

    rows = summary["rows"]
    trend_map: dict[str, dict[str, Any]] = {}
    for row in rows:
        detail = details.get(row["id"], {})
        durations = detail.get("durations", [])
        trend_map[row["id"]] = {
            "summary_24h": row.get("timeout_24h"),
            "summary_36h": row.get("timeout_36h"),
            "detail_count": len(durations),
            "shipment_count": detail.get("shipment_count", 0),
            "pending_count": detail.get("pending_count", 0),
            "invalid_count": detail.get("invalid_count", 0),
            "trend_status": "ready" if detail_status == "uploaded" and durations else detail_status,
            "trends": {
                "24": {"threshold": 24, "points": survival_points(durations, 24)},
                "36": {"threshold": 36, "points": survival_points(durations, 36)},
            },
        }
    source_status = "uploaded" if summary["status"] == "uploaded" and detail_status == "uploaded" else "partial"
    error = detail_error if detail_status != "uploaded" else ""
    return {
        "source_status": source_status,
        "summary_status": summary["status"],
        "detail_status": detail_status,
        "summary_rows": summary["source_rows"],
        "rows": rows,
        "error": error,
        **trace_meta,
    }, trend_map


def clear_tday_source_files(source_dir: Path) -> None:
    source_dir.mkdir(parents=True, exist_ok=True)
    for platform in TDAY_PLATFORMS:
        for role in ROLES:
            (source_dir / f"{platform}_{ROLE_LABELS[role]}.xlsx").unlink(missing_ok=True)


def download_file(url: str, target: Path) -> None:
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request, timeout=180) as response:
        with target.open("wb") as output:
            shutil.copyfileobj(response, output, length=1024 * 1024)
    workbook = load_workbook(target, read_only=True, data_only=True)
    workbook.close()


def install_downloads(downloads: dict[tuple[str, str], Path], source_dir: Path) -> None:
    clear_tday_source_files(source_dir)
    for (platform, role), path in downloads.items():
        target = source_dir / f"{platform}_{ROLE_LABELS[role]}.xlsx"
        shutil.copy2(path, target)


def build_payload(
    source_dir: Path,
    target_day: date,
    mapping: dict[str, dict[str, str]],
    paths: dict[tuple[str, str], Path],
    trace_path: Path | None = None,
) -> dict[str, Any]:
    fallback_snapshot = datetime.now().replace(microsecond=0)
    # ``paths`` 保留在函数签名中以兼容旧调用；T 日现在完全由内网轨迹主链路生成，
    # source_dir 中的任意下载文件只用于可选的单号客户归属补充。
    return build_trace_payload(source_dir, target_day, mapping, trace_path, fallback_snapshot)[0]


def parse_target_day(raw: str | None) -> date:
    if not raw:
        return date.today()
    try:
        return date.fromisoformat(raw)
    except ValueError as error:
        raise ValueError("--date 必须使用 YYYY-MM-DD") from error


def tday_source_freshness(link_file: Path, trace_file: Path) -> list[str]:
    reasons: list[str] = []
    for label, path in (("T日链接清单", link_file), ("T日内网轨迹", trace_file)):
        if not path.is_file():
            reasons.append(f"{label}不存在：{path}")
            continue
        try:
            modified = datetime.fromtimestamp(path.stat().st_mtime).date()
        except OSError:
            reasons.append(f"{label}无法读取修改日期：{path}")
            continue
        if modified != date.today():
            reasons.append(f"{label}最新修改日期为 {modified.isoformat()}，不是今天 {date.today().isoformat()}")
    return reasons


def main() -> int:
    base_dir = Path(__file__).resolve().parent
    data_source_dir = base_dir / "数据源"
    manual_source_dir = data_source_dir / "数据源-手动更新"
    parser = argparse.ArgumentParser(description="生成当天抖音/淘宝 T 日监控数据")
    parser.add_argument("--link-file", type=Path, default=manual_source_dir / "交件链接清单T日.txt")
    parser.add_argument("--source-dir", type=Path, default=data_source_dir / "脚本处理后输出" / "交件T脚本处理后")
    parser.add_argument("--trace-file", type=Path, default=manual_source_dir / "T日内网轨迹.xlsx")
    parser.add_argument("--output", type=Path, default=base_dir / "data" / "dashboard_tday.js")
    parser.add_argument("--date", help="目标 T 日，默认使用本机当天日期")
    parser.add_argument("--download", action="store_true", help="兼容参数：不再在此脚本下载，请先运行独立下载脚本")
    parser.add_argument("--no-download", action="store_true", help="兼容参数：只处理本地内网轨迹（默认行为）")
    args = parser.parse_args()
    freshness_reasons = tday_source_freshness(args.link_file, args.trace_file)
    if freshness_reasons:
        print("T 日数据跳过更新：")
        for reason in freshness_reasons:
            print(f"- {reason}")
        return 0
    target_day = parse_target_day(args.date)
    args.source_dir.mkdir(parents=True, exist_ok=True)

    warnings: list[str] = []
    if args.download:
        warnings.append("下载已拆分为独立脚本，请先运行数据源/下载T日数据.bat")
    try:
        mapping = read_mapping(manual_source_dir)
    except Exception:
        mapping = {}
    payload = build_payload(args.source_dir, target_day, mapping, {}, args.trace_file)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    write_js_payload(args.output, {"t_day": payload}, chunk_name="tday")
    for warning in warnings:
        print(f"提示：{warning}")
    print(f"T 日数据已生成：{target_day.isoformat()} · 状态 {payload['source_status']} · 输出 {args.output}")
    for platform in TDAY_PLATFORMS:
        item = payload["platforms"][platform]
        kpis = item.get("kpis", {})
        print(f"{platform}：内网轨迹 {item['source_status']} · {len(item['rows'])} 条展示 · 待发运 {kpis.get('pending_count', 0)} · 始发退回 {kpis.get('start_return_count', 0)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
